import json
import os
import re

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, range_boundaries


class Evaluator:
    def __init__(self, task_id, target_xlsx, gold_xlsx):
        self.task_id = task_id
        self.target_xlsx = target_xlsx
        self.gold_xlsx = gold_xlsx
        self.failures = []

    def add_failure(self, message):
        self.failures.append(message)

    def require_target(self):
        if not os.path.isfile(self.target_xlsx):
            self.add_failure(f"missing target file: {self.target_xlsx}")
            return False
        return True

    def require_gold(self):
        if not os.path.isfile(self.gold_xlsx):
            self.add_failure(f"missing gold file: {self.gold_xlsx}")
            return False
        return True

    def compare_sheetnames(self):
        target_wb = load_workbook(self.target_xlsx, data_only=False)
        gold_wb = load_workbook(self.gold_xlsx, data_only=False)
        try:
            target_names = list(target_wb.sheetnames)
            gold_names = list(gold_wb.sheetnames)
            if target_names != gold_names:
                self.add_failure(
                    f"sheetnames mismatch: expected {gold_names!r}, got {target_names!r}"
                )
        finally:
            target_wb.close()
            gold_wb.close()

    def _normalize_formula(self, value, worksheet=None, resolve_local_refs=False):
        if not isinstance(value, str) or not value.startswith("="):
            return value
        normalized = re.sub(r"\s+", "", value).replace("$", "")
        if resolve_local_refs and worksheet is not None:
            pattern = r"(?<![A-Za-z0-9_!])([A-Z]{1,3}[1-9][0-9]{0,6})(?![A-Za-z0-9_])"

            def replace_local_ref(match):
                coord = match.group(1)
                cell_value = worksheet[coord].value
                if isinstance(cell_value, str) and not cell_value.startswith("="):
                    return '"' + cell_value.replace('"', '""') + '"'
                return coord

            normalized = re.sub(pattern, replace_local_ref, normalized, flags=re.IGNORECASE)
        return normalized.upper()

    def _normalize_number_format(self, value):
        if not isinstance(value, str):
            return value
        return value.replace("\\", "").strip().lower()

    def _normalize_alignment(self, value):
        if value is None:
            return "general"
        if isinstance(value, str):
            normalized = value.strip().lower()
            if not normalized:
                return "general"
            return normalized
        return value

    def _canonical_rgb(self, value):
        if not isinstance(value, str):
            return value
        value = value.strip().upper()
        if len(value) == 8:
            value = value[-6:]
        if len(value) == 6:
            return value
        return value

    def _canonical_color(self, color):
        normalized = self._normalize_color(color)
        if isinstance(normalized, str):
            return ("rgb", self._canonical_rgb(normalized))
        if (
            isinstance(normalized, tuple)
            and len(normalized) == 3
            and normalized[0] == "theme"
            and normalized[2] in (None, 0, 0.0)
        ):
            if normalized[1] == 0:
                return ("rgb", "FFFFFF")
            if normalized[1] == 1:
                return ("rgb", "000000")
        return normalized

    def _normalize_color(self, color):
        if color is None:
            return None
        color_type = getattr(color, "type", None)
        if color_type == "rgb":
            return getattr(color, "rgb", None)
        if color_type == "indexed":
            return ("indexed", getattr(color, "indexed", None))
        if color_type == "theme":
            return ("theme", getattr(color, "theme", None), getattr(color, "tint", None))
        return getattr(color, "value", None)

    def compare_selected_cells(self, checks, normalize_formulas=False, resolve_local_refs=False):
        target_wb = load_workbook(self.target_xlsx, data_only=False)
        gold_wb = load_workbook(self.gold_xlsx, data_only=False)
        try:
            for sheet_name, coords in checks.items():
                target_ws = target_wb[sheet_name]
                gold_ws = gold_wb[sheet_name]
                for coord in coords:
                    target_value = target_ws[coord].value
                    gold_value = gold_ws[coord].value
                    if normalize_formulas:
                        target_value = self._normalize_formula(
                            target_value,
                            worksheet=target_ws,
                            resolve_local_refs=resolve_local_refs,
                        )
                        gold_value = self._normalize_formula(
                            gold_value,
                            worksheet=gold_ws,
                            resolve_local_refs=resolve_local_refs,
                        )
                    if target_value != gold_value:
                        self.add_failure(
                            f"{sheet_name}!{coord}: expected {gold_value!r}, got {target_value!r}"
                        )
        finally:
            target_wb.close()
            gold_wb.close()

    def compare_selected_cells_data_only(self, checks):
        target_wb = load_workbook(self.target_xlsx, data_only=True)
        gold_wb = load_workbook(self.gold_xlsx, data_only=True)
        try:
            for sheet_name, coords in checks.items():
                target_ws = target_wb[sheet_name]
                gold_ws = gold_wb[sheet_name]
                for coord in coords:
                    target_value = target_ws[coord].value
                    gold_value = gold_ws[coord].value
                    if target_value != gold_value:
                        self.add_failure(
                            f"{sheet_name}!{coord}: expected computed {gold_value!r}, got {target_value!r}"
                        )
        finally:
            target_wb.close()
            gold_wb.close()

    def _evaluate_simple_formula(self, worksheet, value):
        if not isinstance(value, str) or not value.startswith("="):
            return None
        match = re.match(r"^=\s*(SUM|AVERAGE|MAX|MIN)\s*\(\s*([A-Z]{1,3}\d+:[A-Z]{1,3}\d+)\s*\)\s*$", value, re.IGNORECASE)
        if not match:
            return None
        func_name = match.group(1).upper()
        cell_range = match.group(2).upper()
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
        numbers = []
        for row in worksheet.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    numbers.append(cell.value)
        if not numbers:
            return None
        if func_name == "SUM":
            return sum(numbers)
        if func_name == "AVERAGE":
            return sum(numbers) / len(numbers)
        if func_name == "MAX":
            return max(numbers)
        if func_name == "MIN":
            return min(numbers)
        return None

    def _flatten_formula_values(self, values):
        for value in values:
            if isinstance(value, list):
                yield from self._flatten_formula_values(value)
            else:
                yield value

    def _coerce_formula_value(self, value):
        if value is None:
            return 0
        return value

    def _split_reference(self, current_sheet_name, reference):
        if "!" in reference:
            sheet_name, coord = reference.split("!", 1)
            return sheet_name.strip("'"), coord.upper()
        return current_sheet_name, reference.upper()

    def _get_range_values(self, workbook, current_sheet_name, reference, memo, visiting):
        sheet_name, coord_range = self._split_reference(current_sheet_name, reference)
        ws = workbook[sheet_name]
        column_range_match = re.match(r"^([A-Z]{1,3}):([A-Z]{1,3})$", coord_range, re.IGNORECASE)
        if column_range_match:
            min_col = column_index_from_string(column_range_match.group(1).upper())
            max_col = column_index_from_string(column_range_match.group(2).upper())
            min_row = 1
            max_row = ws.max_row
        else:
            min_col, min_row, max_col, max_row = range_boundaries(coord_range)
        values = []
        for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            for cell in row:
                values.append(
                    self._get_computed_cell_value(
                        workbook,
                        sheet_name,
                        cell.coordinate,
                        memo,
                        visiting,
                    )
                )
        return values

    def _replace_formula_tokens(self, expr, pattern, prefix, builder):
        token_map = {}

        def replace(match):
            token = f"__{prefix}_{len(token_map)}__"
            token_map[token] = builder(match.group(0))
            return token

        return pattern.sub(replace, expr), token_map

    def _evaluate_formula(self, workbook, worksheet, value, memo, visiting):
        if not isinstance(value, str) or not value.startswith("="):
            return value
        expr = value[1:].strip().replace("$", "")
        range_pattern = re.compile(
            r"(?<![A-Za-z0-9_\"])"
            r"(?:[A-Za-z_][A-Za-z0-9_]*!)?(?:"
            r"[A-Z]{1,3}[1-9][0-9]{0,6}:[A-Z]{1,3}[1-9][0-9]{0,6}"
            r"|[A-Z]{1,3}:[A-Z]{1,3}"
            r")"
            r"(?![A-Za-z0-9_])",
            re.IGNORECASE,
        )
        cell_pattern = re.compile(
            r"(?<![A-Za-z0-9_\"])"
            r"(?:[A-Za-z_][A-Za-z0-9_]*!)?[A-Z]{1,3}[1-9][0-9]{0,6}"
            r"(?![A-Za-z0-9_])",
            re.IGNORECASE,
        )
        expr, range_tokens = self._replace_formula_tokens(
            expr,
            range_pattern,
            "RANGE",
            lambda ref: f'RANGE_("{ref}")',
        )
        expr, cell_tokens = self._replace_formula_tokens(
            expr,
            cell_pattern,
            "CELL",
            lambda ref: f'CELL_("{ref}")',
        )
        expr = expr.replace("^", "**").replace("<>", "!=")
        expr = re.sub(r"(?<![<>=])=(?![<>=])", "==", expr)
        expr = re.sub(r"\bIF\s*\(", "IF_(", expr, flags=re.IGNORECASE)
        expr = re.sub(r"\bAVERAGEIFS\s*\(", "AVERAGEIFS_(", expr, flags=re.IGNORECASE)
        expr = re.sub(r"\bAVERAGEIF\s*\(", "AVERAGEIF_(", expr, flags=re.IGNORECASE)
        expr = re.sub(r"\bSUMIFS\s*\(", "SUMIFS_(", expr, flags=re.IGNORECASE)
        expr = re.sub(r"\bSUM\s*\(", "SUM_(", expr, flags=re.IGNORECASE)
        expr = re.sub(r"\bAVERAGE\s*\(", "AVERAGE_(", expr, flags=re.IGNORECASE)
        expr = re.sub(r"\bMAX\s*\(", "MAX_(", expr, flags=re.IGNORECASE)
        expr = re.sub(r"\bMIN\s*\(", "MIN_(", expr, flags=re.IGNORECASE)
        expr = re.sub(r"\bSUMIF\s*\(", "SUMIF_(", expr, flags=re.IGNORECASE)
        expr = re.sub(r"\bTRUE\b", "True", expr, flags=re.IGNORECASE)
        expr = re.sub(r"\bFALSE\b", "False", expr, flags=re.IGNORECASE)
        for token_map in (range_tokens, cell_tokens):
            for token, replacement in token_map.items():
                expr = expr.replace(token, replacement)

        def cell_(reference):
            sheet_name, coord = self._split_reference(worksheet.title, reference)
            return self._coerce_formula_value(
                self._get_computed_cell_value(workbook, sheet_name, coord, memo, visiting)
            )

        def range_(reference):
            return [
                self._coerce_formula_value(item)
                for item in self._get_range_values(workbook, worksheet.title, reference, memo, visiting)
            ]

        def normalize_args(args):
            return [item for item in self._flatten_formula_values(args) if item is not None]

        def split_function_args(arg_string):
            parts = []
            current = []
            depth = 0
            in_string = False
            idx = 0
            while idx < len(arg_string):
                char = arg_string[idx]
                if char == '"':
                    in_string = not in_string
                    current.append(char)
                elif not in_string and char == ',' and depth == 0:
                    parts.append(''.join(current).strip())
                    current = []
                else:
                    if not in_string:
                        if char == '(':
                            depth += 1
                        elif char == ')':
                            depth -= 1
                    current.append(char)
                idx += 1
            if current:
                parts.append(''.join(current).strip())
            return parts

        def sumif_(criteria_range, criteria, sum_range=None):
            criteria_values = list(self._flatten_formula_values([criteria_range]))
            if sum_range is None:
                sum_values = criteria_values
            else:
                sum_values = list(self._flatten_formula_values([sum_range]))
            total = 0
            for idx, candidate in enumerate(criteria_values):
                if candidate == criteria:
                    if idx < len(sum_values) and isinstance(sum_values[idx], (int, float)):
                        total += sum_values[idx]
            return total

        def averageif_(criteria_range, criteria, average_range=None):
            criteria_values = list(self._flatten_formula_values([criteria_range]))
            if average_range is None:
                average_values = criteria_values
            else:
                average_values = list(self._flatten_formula_values([average_range]))
            total = 0
            count = 0
            for idx, candidate in enumerate(criteria_values):
                if candidate == criteria and idx < len(average_values):
                    value = average_values[idx]
                    if isinstance(value, (int, float)):
                        total += value
                        count += 1
            return total / count if count else 0

        def sumifs_(sum_range, *criteria_pairs):
            sum_values = list(self._flatten_formula_values([sum_range]))
            if len(criteria_pairs) % 2 != 0:
                return 0
            parsed_pairs = []
            for idx in range(0, len(criteria_pairs), 2):
                criteria_values = list(self._flatten_formula_values([criteria_pairs[idx]]))
                criteria = criteria_pairs[idx + 1]
                parsed_pairs.append((criteria_values, criteria))
            total = 0
            for idx, value in enumerate(sum_values):
                matches = True
                for criteria_values, criteria in parsed_pairs:
                    candidate = criteria_values[idx] if idx < len(criteria_values) else None
                    if candidate != criteria:
                        matches = False
                        break
                if matches and isinstance(value, (int, float)):
                    total += value
            return total

        def averageifs_(average_range, *criteria_pairs):
            average_values = list(self._flatten_formula_values([average_range]))
            if len(criteria_pairs) % 2 != 0:
                return 0
            parsed_pairs = []
            for idx in range(0, len(criteria_pairs), 2):
                criteria_values = list(self._flatten_formula_values([criteria_pairs[idx]]))
                criteria = criteria_pairs[idx + 1]
                parsed_pairs.append((criteria_values, criteria))
            total = 0
            count = 0
            for idx, value in enumerate(average_values):
                matches = True
                for criteria_values, criteria in parsed_pairs:
                    candidate = criteria_values[idx] if idx < len(criteria_values) else None
                    if candidate != criteria:
                        matches = False
                        break
                if matches and isinstance(value, (int, float)):
                    total += value
                    count += 1
            return total / count if count else 0

        namespace = {
            "CELL_": cell_,
            "RANGE_": range_,
            "IF_": lambda cond, when_true, when_false: when_true if cond else when_false,
            "SUM_": lambda *args: sum(normalize_args(args)),
            "AVERAGE_": lambda *args: (
                sum(normalize_args(args)) / len(normalize_args(args))
                if normalize_args(args)
                else 0
            ),
            "MAX_": lambda *args: max(normalize_args(args)) if normalize_args(args) else 0,
            "MIN_": lambda *args: min(normalize_args(args)) if normalize_args(args) else 0,
            "AVERAGEIF_": averageif_,
            "AVERAGEIFS_": averageifs_,
            "SUMIF_": sumif_,
            "SUMIFS_": sumifs_,
        }

        def eval_expression(inner_expr):
            stripped = inner_expr.strip()
            if stripped.upper().startswith("IF_(") and stripped.endswith(")"):
                args = split_function_args(stripped[4:-1])
                if len(args) == 3:
                    condition = eval_expression(args[0])
                    branch = args[1] if condition else args[2]
                    return eval_expression(branch)
            return eval(stripped, {"__builtins__": {}}, namespace)
        try:
            return eval_expression(expr)
        except Exception:
            return None

    def _get_computed_cell_value(self, workbook, sheet_name, coord, memo, visiting):
        key = (sheet_name, coord)
        if key in memo:
            return memo[key]
        if key in visiting:
            return None
        visiting.add(key)
        ws = workbook[sheet_name]
        raw_value = ws[coord].value
        if isinstance(raw_value, str) and raw_value.startswith("="):
            value = self._evaluate_formula(workbook, ws, raw_value, memo, visiting)
        else:
            value = raw_value
        memo[key] = value
        visiting.remove(key)
        return value

    def compare_formula_cells_by_value(self, checks, tolerance=1e-9):
        target_formula_wb = load_workbook(self.target_xlsx, data_only=False)
        try:
            memo = {}
            for sheet_name, coord_map in checks.items():
                formula_ws = target_formula_wb[sheet_name]
                for coord, expected_value in coord_map.items():
                    formula_value = formula_ws[coord].value
                    if not isinstance(formula_value, str) or not formula_value.startswith("="):
                        self.add_failure(f"{sheet_name}!{coord}: expected a formula cell, got {formula_value!r}")
                        continue
                    actual_value = self._get_computed_cell_value(
                        target_formula_wb,
                        sheet_name,
                        coord,
                        memo,
                        set(),
                    )
                    if actual_value is None:
                        actual_value = self._evaluate_simple_formula(formula_ws, formula_value)
                    if actual_value is None:
                        self.add_failure(f"{sheet_name}!{coord}: formula result is unavailable")
                        continue
                    if isinstance(expected_value, (int, float)) and isinstance(actual_value, (int, float)):
                        if abs(actual_value - expected_value) > tolerance:
                            self.add_failure(
                                f"{sheet_name}!{coord}: expected computed {expected_value!r}, got {actual_value!r}"
                            )
                    elif actual_value != expected_value:
                        self.add_failure(
                            f"{sheet_name}!{coord}: expected computed {expected_value!r}, got {actual_value!r}"
                        )
        finally:
            target_formula_wb.close()

    def compare_protected_cells(self, checks):
        self.compare_selected_cells(checks)

    def compare_sheet_dimension(self, sheet_name):
        target_wb = load_workbook(self.target_xlsx, data_only=False)
        gold_wb = load_workbook(self.gold_xlsx, data_only=False)
        try:
            target_ws = target_wb[sheet_name]
            gold_ws = gold_wb[sheet_name]
            if target_ws.max_row != gold_ws.max_row:
                self.add_failure(
                    f"{sheet_name} max_row mismatch: expected {gold_ws.max_row}, got {target_ws.max_row}"
                )
            if target_ws.max_column != gold_ws.max_column:
                self.add_failure(
                    f"{sheet_name} max_column mismatch: expected {gold_ws.max_column}, got {target_ws.max_column}"
                )
        finally:
            target_wb.close()
            gold_wb.close()

    def compare_cell_styles(
        self,
        style_checks,
        normalize_number_formats=False,
        compare_font_color=False,
        compare_alignment=False,
        compare_borders=False,
    ):
        target_wb = load_workbook(self.target_xlsx, data_only=False)
        gold_wb = load_workbook(self.gold_xlsx, data_only=False)
        try:
            for sheet_name, coords in style_checks.items():
                target_ws = target_wb[sheet_name]
                gold_ws = gold_wb[sheet_name]
                for coord in coords:
                    target_cell = target_ws[coord]
                    gold_cell = gold_ws[coord]
                    target_number_format = target_cell.number_format
                    gold_number_format = gold_cell.number_format
                    if normalize_number_formats:
                        target_number_format = self._normalize_number_format(target_number_format)
                        gold_number_format = self._normalize_number_format(gold_number_format)
                    if target_number_format != gold_number_format:
                        self.add_failure(
                            f"{sheet_name}!{coord} number_format mismatch: expected "
                            f"{gold_cell.number_format!r}, got {target_cell.number_format!r}"
                        )
                    if target_cell.font.bold != gold_cell.font.bold:
                        self.add_failure(
                            f"{sheet_name}!{coord} bold mismatch: expected "
                            f"{gold_cell.font.bold!r}, got {target_cell.font.bold!r}"
                        )
                    if compare_font_color:
                        target_font_color = self._canonical_color(target_cell.font.color)
                        gold_font_color = self._canonical_color(gold_cell.font.color)
                        if target_font_color != gold_font_color:
                            self.add_failure(
                                f"{sheet_name}!{coord} font color mismatch: expected "
                                f"{gold_font_color!r}, got {target_font_color!r}"
                            )
                    if compare_alignment:
                        target_alignment = self._normalize_alignment(target_cell.alignment.horizontal)
                        gold_alignment = self._normalize_alignment(gold_cell.alignment.horizontal)
                        if target_alignment != gold_alignment:
                            self.add_failure(
                                f"{sheet_name}!{coord} horizontal alignment mismatch: expected "
                                f"{gold_cell.alignment.horizontal!r}, got {target_cell.alignment.horizontal!r}"
                            )
                    if compare_borders:
                        for side_name in ("left", "right", "top", "bottom"):
                            target_side = getattr(target_cell.border, side_name)
                            gold_side = getattr(gold_cell.border, side_name)
                            if target_side.style != gold_side.style:
                                self.add_failure(
                                    f"{sheet_name}!{coord} {side_name} border mismatch: expected "
                                    f"{gold_side.style!r}, got {target_side.style!r}"
                                )
                    target_fill_color = self._canonical_color(target_cell.fill.fgColor)
                    gold_fill_color = self._canonical_color(gold_cell.fill.fgColor)
                    if (
                        target_cell.fill.fill_type != gold_cell.fill.fill_type
                        or target_fill_color != gold_fill_color
                    ):
                        self.add_failure(f"{sheet_name}!{coord} fill mismatch")
        finally:
            target_wb.close()
            gold_wb.close()

    def compare_sheet_view_settings(self, checks):
        target_wb = load_workbook(self.target_xlsx, data_only=False)
        gold_wb = load_workbook(self.gold_xlsx, data_only=False)
        try:
            for sheet_name, options in checks.items():
                target_ws = target_wb[sheet_name]
                gold_ws = gold_wb[sheet_name]
                if options.get("freeze_panes"):
                    target_freeze = target_ws.freeze_panes
                    gold_freeze = gold_ws.freeze_panes
                    target_coord = None if target_freeze is None else getattr(target_freeze, "coordinate", target_freeze)
                    gold_coord = None if gold_freeze is None else getattr(gold_freeze, "coordinate", gold_freeze)
                    if target_coord != gold_coord:
                        self.add_failure(
                            f"{sheet_name} freeze_panes mismatch: expected {gold_coord!r}, got {target_coord!r}"
                        )
                if options.get("auto_filter"):
                    if target_ws.auto_filter.ref != gold_ws.auto_filter.ref:
                        self.add_failure(
                            f"{sheet_name} auto_filter mismatch: expected {gold_ws.auto_filter.ref!r}, "
                            f"got {target_ws.auto_filter.ref!r}"
                        )
        finally:
            target_wb.close()
            gold_wb.close()

    def compare_column_dimensions(self, checks, compare_hidden=False, tolerance=1e-9):
        target_wb = load_workbook(self.target_xlsx, data_only=False)
        gold_wb = load_workbook(self.gold_xlsx, data_only=False)
        try:
            for sheet_name, columns in checks.items():
                target_ws = target_wb[sheet_name]
                gold_ws = gold_wb[sheet_name]
                for column_name in columns:
                    target_dim = target_ws.column_dimensions[column_name]
                    gold_dim = gold_ws.column_dimensions[column_name]
                    target_width = target_dim.width
                    gold_width = gold_dim.width
                    if target_width is None or gold_width is None:
                        if target_width != gold_width:
                            self.add_failure(
                                f"{sheet_name}!{column_name} width mismatch: expected {gold_width!r}, got {target_width!r}"
                            )
                    elif abs(target_width - gold_width) > tolerance:
                        self.add_failure(
                            f"{sheet_name}!{column_name} width mismatch: expected {gold_width!r}, got {target_width!r}"
                        )
                    if compare_hidden and target_dim.hidden != gold_dim.hidden:
                        self.add_failure(
                            f"{sheet_name}!{column_name} hidden mismatch: expected {gold_dim.hidden!r}, got {target_dim.hidden!r}"
                        )
        finally:
            target_wb.close()
            gold_wb.close()

    def compare_row_dimensions(self, checks, tolerance=1e-9):
        target_wb = load_workbook(self.target_xlsx, data_only=False)
        gold_wb = load_workbook(self.gold_xlsx, data_only=False)
        try:
            for sheet_name, rows in checks.items():
                target_ws = target_wb[sheet_name]
                gold_ws = gold_wb[sheet_name]
                for row_num in rows:
                    target_height = target_ws.row_dimensions[row_num].height
                    gold_height = gold_ws.row_dimensions[row_num].height
                    if target_height is None or gold_height is None:
                        if target_height != gold_height:
                            self.add_failure(
                                f"{sheet_name}!row{row_num} height mismatch: expected {gold_height!r}, got {target_height!r}"
                            )
                    elif abs(target_height - gold_height) > tolerance:
                        self.add_failure(
                            f"{sheet_name}!row{row_num} height mismatch: expected {gold_height!r}, got {target_height!r}"
                        )
        finally:
            target_wb.close()
            gold_wb.close()

    def write_result(self, output_path, extra=None):
        parent = os.path.dirname(output_path)
        if parent and not os.path.exists(parent):
            os.makedirs(parent, exist_ok=True)
        result = {
            "task_id": self.task_id,
            "passed": len(self.failures) == 0,
            "failure_count": len(self.failures),
            "failures": list(self.failures),
        }
        if extra:
            result.update(extra)
        with open(output_path, "w", encoding="utf-8") as f:
            json.dump(result, f, ensure_ascii=False, indent=2)
