import argparse
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Font


DATA_ROW_COUNT = 220
CATEGORY_CANONICAL = ["Travel", "Software", "Office Supplies", "Meals", "Marketing"]
CATEGORY_VARIANTS = {
    "Travel": ["travel", "Travel", " TRAVEL ", "travel "],
    "Software": ["software", "Software", " SOFTWARE", "software "],
    "Office Supplies": ["office supplies", "Office Supplies", " OFFICE SUPPLIES ", "office supplies "],
    "Meals": ["meals", "Meals", " MEALS ", "meals "],
    "Marketing": ["marketing", "Marketing", " MARKETING ", "marketing "],
}
OWNER_CANONICAL = ["Ava Chen", "Leo Wang", "Mia Liu", "Noah Sun", "Ivy Zhao"]
OWNER_VARIANTS = {
    "Ava Chen": ["Ava", "ava chen", "AVA CHEN", "Ava Chen "],
    "Leo Wang": ["Leo", "leo wang", "LEO WANG", " Leo Wang"],
    "Mia Liu": ["Mia", "mia liu", "MIA LIU", "Mia Liu "],
    "Noah Sun": ["Noah", "noah sun", "NOAH SUN", " Noah Sun"],
    "Ivy Zhao": ["Ivy", "ivy zhao", "IVY ZHAO", "Ivy Zhao "],
}
COST_CENTER_CANONICAL = ["CC-100", "CC-200", "CC-300", "CC-400"]


def canonical_rows():
    rows = []
    for i in range(1, DATA_ROW_COUNT + 1):
        month = (i - 1) % 6 + 1
        day = (i * 3 - 1) % 28 + 1
        date_value = f"2026-{month:02d}-{day:02d}"
        category = CATEGORY_CANONICAL[(i - 1) % len(CATEGORY_CANONICAL)]
        owner = OWNER_CANONICAL[(i - 1) % len(OWNER_CANONICAL)]
        cost_center = COST_CENTER_CANONICAL[(i - 1) % len(COST_CENTER_CANONICAL)]
        amount = round(25 + i * 4.35, 2)
        rows.append(
            {
                "date": date_value,
                "category": category,
                "amount": amount,
                "owner": owner,
                "cost_center": cost_center,
            }
        )
    return rows


def messy_rows():
    rows = []
    for index, row in enumerate(canonical_rows(), start=1):
        date_variants = [
            row["date"].replace("-", "/"),
            row["date"],
            row["date"].replace("-", "."),
            f" {row['date']} ",
        ]
        amount_variants = [
            f"{row['amount']:.2f}",
            f"{row['amount']:,.2f}",
            f"${row['amount']:.2f}",
            f" {row['amount']:.2f} ",
        ]
        rows.append(
            [
                date_variants[index % len(date_variants)],
                CATEGORY_VARIANTS[row["category"]][index % len(CATEGORY_VARIANTS[row["category"]])],
                amount_variants[index % len(amount_variants)],
                OWNER_VARIANTS[row["owner"]][index % len(OWNER_VARIANTS[row["owner"]])],
                row["cost_center"].lower() if index % 2 == 0 else row["cost_center"],
            ]
        )
    return rows


def build_workbook(output_file):
    wb = Workbook()
    ws = wb.active
    ws.title = "Raw"
    ws["A1"] = "Expense export"
    headers = ["Date", "Category", "Amount", "Owner", "Cost Center"]
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=2, column=col_idx, value=header)
    for row_index, row in enumerate(messy_rows(), start=3):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_index, column=col_idx, value=value)
    ws["G2"] = "Ignore side notes"
    ws["G50"] = "reconciliation pending"
    ws["H160"] = "legacy export"
    for cell in ws[2]:
        cell.font = Font(bold=True)
    wb.save(output_file)
    wb.close()


def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    output_file = os.path.join(args.output_path, "expense_raw.xlsx")
    build_workbook(output_file)
    with open(os.path.join(args.output_path, "task_metadata.json"), "w", encoding="utf-8") as f:
        json.dump(
            {
                "task_id": "data_cleaning/task_002_normalize_expense_export",
                "input_file": "expense_raw.xlsx",
                "target_file": "expense_clean.xlsx",
            },
            f,
            ensure_ascii=False,
            indent=2,
        )
    print(os.path.abspath(args.output_path))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Create workspace for data cleaning task")
    parser.add_argument("output_path", help="Directory to create workspace in")
    main(parser.parse_args())
