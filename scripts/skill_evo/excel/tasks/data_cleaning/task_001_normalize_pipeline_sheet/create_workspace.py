import argparse
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Font


DATA_ROW_COUNT = 220
STAGE_CANONICAL = ["Proposal", "Closed Won", "Qualified", "Negotiation"]
STAGE_VARIANTS = {
    "Proposal": ["proposal", "Proposal", " PROPOSAL ", "proposal "],
    "Closed Won": ["closed won", "Closed Won", "CLOSED WON", " closed won "],
    "Qualified": ["qualified", "Qualified", "QUALIFIED ", " qualified"],
    "Negotiation": ["negotiation", "Negotiation", "NEGOTIATION", " negotiation "],
}
OWNER_CANONICAL = ["Ava Chen", "Leo Wang", "Mia Liu", "Noah Sun", "Ivy Zhao"]
OWNER_VARIANTS = {
    "Ava Chen": ["Ava", "ava chen", "AVA CHEN", "Ava Chen "],
    "Leo Wang": ["Leo", "leo wang", "LEO WANG", " Leo Wang"],
    "Mia Liu": ["Mia", "mia liu", "MIA LIU", "Mia Liu "],
    "Noah Sun": ["Noah", "noah sun", "NOAH SUN", " Noah Sun"],
    "Ivy Zhao": ["Ivy", "ivy zhao", "IVY ZHAO", "Ivy Zhao "],
}
REGION_CANONICAL = ["North", "South", "West", "East"]
REGION_VARIANTS = {
    "North": ["north", "North", "NORTH ", " north"],
    "South": ["south", "South", "SOUTH ", " south"],
    "West": ["west", "West", "WEST ", " west"],
    "East": ["east", "East", "EAST ", " east"],
}


def canonical_rows():
    rows = []
    for i in range(1, DATA_ROW_COUNT + 1):
        stage = STAGE_CANONICAL[(i - 1) % len(STAGE_CANONICAL)]
        owner = OWNER_CANONICAL[(i - 1) % len(OWNER_CANONICAL)]
        region = REGION_CANONICAL[(i - 1) % len(REGION_CANONICAL)]
        amount = 500 + i * 17
        rows.append(
            {
                "deal_id": f"D-{i:04d}",
                "stage": stage,
                "amount": amount,
                "owner": owner,
                "region": region,
            }
        )
    return rows


def messy_rows():
    rows = []
    for index, row in enumerate(canonical_rows(), start=1):
        stage_value = STAGE_VARIANTS[row["stage"]][index % len(STAGE_VARIANTS[row["stage"]])]
        owner_value = OWNER_VARIANTS[row["owner"]][index % len(OWNER_VARIANTS[row["owner"]])]
        region_value = REGION_VARIANTS[row["region"]][index % len(REGION_VARIANTS[row["region"]])]
        amount_value = [
            str(row["amount"]),
            f"{row['amount']:,}",
            f" {row['amount']} ",
            f"${row['amount']}",
        ][index % 4]
        rows.append([row["deal_id"], stage_value, amount_value, owner_value, region_value])
    return rows


def build_workbook(output_file):
    wb = Workbook()
    ws = wb.active
    ws.title = "Raw"
    ws["A1"] = "Pipeline export generated from CRM extract"
    ws["G1"] = "Notes"
    ws["G2"] = "Ignore columns G:H and blank spacer rows."
    headers = ["Deal ID", "Stage", "Amount", "Owner", "Region"]
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=3, column=col_idx, value=header)
    for row_index, row in enumerate(messy_rows(), start=4):
        if row_index in (40, 120, 180):
            ws.cell(row=row_index, column=7, value="manual check")
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_index, column=col_idx, value=value)
    ws["H10"] = "staging area"
    ws["H150"] = "side comment"
    for cell in ws[3]:
        cell.font = Font(bold=True)
    wb.save(output_file)
    wb.close()


def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    output_file = os.path.join(args.output_path, "pipeline_raw.xlsx")
    build_workbook(output_file)
    with open(os.path.join(args.output_path, "task_metadata.json"), "w", encoding="utf-8") as f:
        json.dump(
            {
                "task_id": "data_cleaning/task_001_normalize_pipeline_sheet",
                "input_file": "pipeline_raw.xlsx",
                "target_file": "pipeline_clean.xlsx",
            },
            f,
            ensure_ascii=False,
            indent=2,
        )
    print(os.path.abspath(args.output_path))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Create workspace for data cleaning task")
    parser.add_argument("output_path", help="Directory to create workspace in")
    main(parser.parse_args())
