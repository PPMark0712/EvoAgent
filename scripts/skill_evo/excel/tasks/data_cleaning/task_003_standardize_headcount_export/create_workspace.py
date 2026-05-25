import argparse
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Font


DATA_ROW_COUNT = 220
DEPARTMENT_CANONICAL = ["Engineering", "Sales", "Finance", "Operations", "People"]
DEPARTMENT_VARIANTS = {
    "Engineering": ["engineering", "ENGINEERING", "Engineering ", " engineering"],
    "Sales": ["sales", "SALES", "Sales ", " sales"],
    "Finance": ["finance", "FINANCE", "Finance ", " finance"],
    "Operations": ["operations", "OPERATIONS", "Operations ", " operations"],
    "People": ["people", "PEOPLE", "People ", " people"],
}
LOCATION_CANONICAL = ["Beijing", "Shanghai", "Shenzhen", "Singapore"]
LOCATION_VARIANTS = {
    "Beijing": ["beijing", "BEIJING", "Beijing ", " beijing"],
    "Shanghai": ["shanghai", "SHANGHAI", "Shanghai ", " shanghai"],
    "Shenzhen": ["shenzhen", "SHENZHEN", "Shenzhen ", " shenzhen"],
    "Singapore": ["singapore", "SINGAPORE", "Singapore ", " singapore"],
}
STATUS_CANONICAL = ["Active", "Leave", "Inactive"]
STATUS_VARIANTS = {
    "Active": ["active", "ACTIVE", "Active ", " active"],
    "Leave": ["leave", "LEAVE", "Leave ", " leave"],
    "Inactive": ["inactive", "INACTIVE", "Inactive ", " inactive"],
}
MANAGER_CANONICAL = ["Nina Gu", "Oscar Hu", "Pia Lin", "Ray Xu"]
MANAGER_VARIANTS = {
    "Nina Gu": ["Nina", "nina gu", "NINA GU", "Nina Gu "],
    "Oscar Hu": ["Oscar", "oscar hu", "OSCAR HU", " Oscar Hu"],
    "Pia Lin": ["Pia", "pia lin", "PIA LIN", " Pia Lin"],
    "Ray Xu": ["Ray", "ray xu", "RAY XU", "Ray Xu "],
}


def canonical_rows():
    rows = []
    for i in range(1, DATA_ROW_COUNT + 1):
        rows.append(
            {
                "employee_id": f"E-{i:04d}",
                "department": DEPARTMENT_CANONICAL[(i - 1) % len(DEPARTMENT_CANONICAL)],
                "location": LOCATION_CANONICAL[(i - 1) % len(LOCATION_CANONICAL)],
                "status": STATUS_CANONICAL[(i - 1) % len(STATUS_CANONICAL)],
                "manager": MANAGER_CANONICAL[(i - 1) % len(MANAGER_CANONICAL)],
            }
        )
    return rows


def messy_rows():
    rows = []
    for index, row in enumerate(canonical_rows(), start=1):
        rows.append(
            [
                row["employee_id"],
                DEPARTMENT_VARIANTS[row["department"]][index % len(DEPARTMENT_VARIANTS[row["department"]])],
                LOCATION_VARIANTS[row["location"]][index % len(LOCATION_VARIANTS[row["location"]])],
                STATUS_VARIANTS[row["status"]][index % len(STATUS_VARIANTS[row["status"]])],
                MANAGER_VARIANTS[row["manager"]][index % len(MANAGER_VARIANTS[row["manager"]])],
            ]
        )
    return rows


def build_workbook(output_file):
    wb = Workbook()
    ws = wb.active
    ws.title = "Raw"
    ws["A1"] = "Headcount export - latest snapshot"
    headers = ["Employee ID", "Department", "Location", "Status", "Manager"]
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=3, column=col_idx, value=header)
    for row_index, row in enumerate(messy_rows(), start=4):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_index, column=col_idx, value=value)
    ws["G3"] = "Notes"
    ws["G4"] = "Ignore side column"
    ws["H120"] = "manual roster merge"
    for cell in ws[3]:
        cell.font = Font(bold=True)
    wb.save(output_file)
    wb.close()


def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    output_file = os.path.join(args.output_path, "roster_raw.xlsx")
    build_workbook(output_file)
    with open(os.path.join(args.output_path, "task_metadata.json"), "w", encoding="utf-8") as f:
        json.dump(
            {
                "task_id": "data_cleaning/task_003_standardize_headcount_export",
                "input_file": "roster_raw.xlsx",
                "target_file": "roster_clean.xlsx",
            },
            f,
            ensure_ascii=False,
            indent=2,
        )
    print(os.path.abspath(args.output_path))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Create workspace for data cleaning task")
    parser.add_argument("output_path", help="Directory to create workspace in")
    main(parser.parse_args())
