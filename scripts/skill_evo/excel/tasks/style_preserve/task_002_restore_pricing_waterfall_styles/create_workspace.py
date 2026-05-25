import argparse
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


DATA_ROW_COUNT = 180
GUIDE_ROW_COUNT = 18
HEADER_FILL = PatternFill('solid', fgColor='5B9BD5')
HEADER_FONT = Font(bold=True, color='FFFFFFFF')
DATA_FONT = Font(color='FF000000')
TARGET_FILL = PatternFill('solid', fgColor='DDEBF7')


def build_rows(count, start_index=1):
    rows = []
    for idx in range(start_index, start_index + count):
        product = f'Product {idx:03d}'
        region = ['North', 'South', 'East', 'West'][(idx - 1) % 4]
        list_price = 120 + (idx % 11) * 7
        discount = round(0.04 + (idx % 5) * 0.02, 4)
        units = 80 + (idx % 12) * 5
        unit_cost = 54 + (idx % 7) * 3
        rows.append((product, region, list_price, discount, None, units, unit_cost, None, None))
    return rows


def write_formulas(ws, row_num, broken):
    if broken:
        ws[f'E{row_num}'] = f'=C{row_num}*(1+D{row_num})'
        ws[f'H{row_num}'] = f'=E{row_num}+F{row_num}'
        ws[f'I{row_num}'] = f'=IF(G{row_num}=0,0,H{row_num}/G{row_num})'
    else:
        ws[f'E{row_num}'] = f'=C{row_num}*(1-D{row_num})'
        ws[f'H{row_num}'] = f'=E{row_num}*F{row_num}'
        ws[f'I{row_num}'] = f'=IF(H{row_num}=0,0,(H{row_num}-F{row_num}*G{row_num})/H{row_num})'


def apply_full_styles(ws, row_count):
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center')
    for row_num in range(2, row_count + 2):
        for col in ('C', 'E', 'G', 'H'):
            ws[f'{col}{row_num}'].number_format = '$#,##0.00'
            ws[f'{col}{row_num}'].font = DATA_FONT
        ws[f'F{row_num}'].number_format = '#,##0'
        ws[f'F{row_num}'].font = DATA_FONT
        for col in ('D', 'I'):
            ws[f'{col}{row_num}'].number_format = '0.0%'
            ws[f'{col}{row_num}'].font = DATA_FONT
        for col in ('E', 'H', 'I'):
            ws[f'{col}{row_num}'].fill = TARGET_FILL
            ws[f'{col}{row_num}'].alignment = Alignment(horizontal='center')


def corrupt_target_styles(ws, row_count):
    for row_num in range(2, row_count + 2):
        for col in ('E', 'H', 'I'):
            ws[f'{col}{row_num}'].number_format = 'General'
            ws[f'{col}{row_num}'].fill = PatternFill()
            ws[f'{col}{row_num}'].font = Font(color='FF0000FF')
            ws[f'{col}{row_num}'].alignment = Alignment(horizontal='left')


def build_workbook(output_file):
    wb = Workbook()
    reference = wb.active
    reference.title = 'StyleGuide'
    target = wb.create_sheet('Waterfall')
    reference.append(['Product', 'Region', 'List Price', 'Discount %', 'Net Price', 'Units', 'Unit Cost', 'Revenue', 'Gross Margin %'])
    target.append(['Product', 'Region', 'List Price', 'Discount %', 'Net Price', 'Units', 'Unit Cost', 'Revenue', 'Gross Margin %'])
    for row in build_rows(GUIDE_ROW_COUNT, 1):
        reference.append(list(row))
    for row in build_rows(DATA_ROW_COUNT, GUIDE_ROW_COUNT + 1):
        target.append(list(row))
    for row_num in range(2, GUIDE_ROW_COUNT + 2):
        write_formulas(reference, row_num, broken=False)
    for row_num in range(2, DATA_ROW_COUNT + 2):
        write_formulas(target, row_num, broken=True)
    apply_full_styles(reference, GUIDE_ROW_COUNT)
    apply_full_styles(target, DATA_ROW_COUNT)
    corrupt_target_styles(target, DATA_ROW_COUNT)
    wb.save(output_file)
    wb.close()


def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    output_file = os.path.join(args.output_path, 'pricing_waterfall_style.xlsx')
    build_workbook(output_file)
    with open(os.path.join(args.output_path, 'task_metadata.json'), 'w', encoding='utf-8') as f:
        json.dump({'task_id': 'style_preserve/task_002_restore_pricing_waterfall_styles', 'target_file': 'pricing_waterfall_style.xlsx'}, f, ensure_ascii=False, indent=2)
    print(os.path.abspath(args.output_path))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Create workspace for style preserve task')
    parser.add_argument('output_path', help='Directory to create workspace in')
    main(parser.parse_args())
