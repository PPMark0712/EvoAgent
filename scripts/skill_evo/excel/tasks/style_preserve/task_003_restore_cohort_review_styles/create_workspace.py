import argparse
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


DATA_ROW_COUNT = 180
GUIDE_ROW_COUNT = 18
HEADER_FILL = PatternFill('solid', fgColor='70AD47')
HEADER_FONT = Font(bold=True, color='FFFFFFFF')
DATA_FONT = Font(color='FF000000')
TARGET_FILL = PatternFill('solid', fgColor='E2F0D9')


def build_rows(count, start_index=1):
    rows = []
    for idx in range(start_index, start_index + count):
        cohort = f'Cohort {idx:03d}'
        region = ['North', 'South', 'East', 'West'][(idx - 1) % 4]
        start_mrr = 9400 + idx * 62
        expansion = round(0.04 + (idx % 4) * 0.015, 4)
        churn = round(0.01 + (idx % 5) * 0.01, 4)
        rows.append((cohort, region, start_mrr, expansion, churn))
    return rows


def write_formulas(ws, row_num, broken):
    if broken:
        ws[f'F{row_num}'] = f'=C{row_num}*(1+D{row_num}+E{row_num})'
        ws[f'G{row_num}'] = f'=IF(C{row_num}=0,0,F{row_num}/C{row_num}-1)'
    else:
        ws[f'F{row_num}'] = f'=C{row_num}*(1+D{row_num}-E{row_num})'
        ws[f'G{row_num}'] = f'=IF(C{row_num}=0,0,F{row_num}/C{row_num})'


def apply_full_styles(ws, row_count):
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center')
    for row_num in range(2, row_count + 2):
        for col in ('C', 'F'):
            ws[f'{col}{row_num}'].number_format = '$#,##0'
            ws[f'{col}{row_num}'].font = DATA_FONT
        for col in ('D', 'E', 'G'):
            ws[f'{col}{row_num}'].number_format = '0.0%'
            ws[f'{col}{row_num}'].font = DATA_FONT
        for col in ('F', 'G'):
            ws[f'{col}{row_num}'].fill = TARGET_FILL
            ws[f'{col}{row_num}'].alignment = Alignment(horizontal='center')


def corrupt_target_styles(ws, row_count):
    for row_num in range(2, row_count + 2):
        for col in ('F', 'G'):
            ws[f'{col}{row_num}'].number_format = 'General'
            ws[f'{col}{row_num}'].fill = PatternFill()
            ws[f'{col}{row_num}'].font = Font(color='FFFF0000')
            ws[f'{col}{row_num}'].alignment = Alignment(horizontal='left')


def build_workbook(output_file):
    wb = Workbook()
    reference = wb.active
    reference.title = 'Template'
    target = wb.create_sheet('Cohorts')
    reference.append(['Cohort', 'Region', 'Start MRR', 'Expansion %', 'Churn %', 'End MRR', 'Net Retention %'])
    target.append(['Cohort', 'Region', 'Start MRR', 'Expansion %', 'Churn %', 'End MRR', 'Net Retention %'])
    for row in build_rows(GUIDE_ROW_COUNT, 1):
        reference.append(list(row))
    for row in build_rows(DATA_ROW_COUNT, GUIDE_ROW_COUNT + 1):
        target.append(list(row))
    for row_num in range(2, GUIDE_ROW_COUNT + 2):
        write_formulas(reference, row_num, broken=False)
    for row_num in range(2, DATA_ROW_COUNT + 2):
        write_formulas(target, row_num, broken=True)
    apply_full_styles(reference, GUIDE_ROW_COUNT)
    apply_full_styles(target, DATA_ROW_COUNT)
    corrupt_target_styles(target, DATA_ROW_COUNT)
    wb.save(output_file)
    wb.close()


def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    output_file = os.path.join(args.output_path, 'cohort_review_style.xlsx')
    build_workbook(output_file)
    with open(os.path.join(args.output_path, 'task_metadata.json'), 'w', encoding='utf-8') as f:
        json.dump({'task_id': 'style_preserve/task_003_restore_cohort_review_styles', 'target_file': 'cohort_review_style.xlsx'}, f, ensure_ascii=False, indent=2)
    print(os.path.abspath(args.output_path))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Create workspace for style preserve task')
    parser.add_argument('output_path', help='Directory to create workspace in')
    main(parser.parse_args())
