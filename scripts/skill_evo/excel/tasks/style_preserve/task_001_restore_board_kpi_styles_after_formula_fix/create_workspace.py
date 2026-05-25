import argparse
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


DATA_ROW_COUNT = 180
GUIDE_ROW_COUNT = 18
HEADER_FILL = PatternFill('solid', fgColor='203864')
HEADER_FONT = Font(bold=True, color='FFFFFFFF')
DATA_FONT = Font(color='FF000000')
TARGET_FILL = PatternFill('solid', fgColor='FCE4D6')
LOCAL_FILL = PatternFill('solid', fgColor='E2F0D9')
THIN_SIDE = Side(style='thin', color='D9D9D9')
TARGET_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
COLUMN_WIDTHS = {
    'A': 14.0,
    'B': 12.0,
    'C': 13.0,
    'D': 13.0,
    'E': 14.0,
    'F': 12.0,
    'G': 12.5,
    'H': 13.5,
}


def build_rows(count, start_index=1):
    rows = []
    for idx in range(start_index, start_index + count):
        team = f'Team {idx:03d}'
        region = ['North', 'South', 'East', 'West'][(idx - 1) % 4]
        q1 = 0 if idx % 47 == 0 else 8200 + idx * 71
        q2 = q1 + 460 + (idx % 6) * 85
        renewal = round(0.78 + (idx % 5) * 0.025, 4)
        expansion = round(0.06 + (idx % 4) * 0.015, 4)
        rows.append((team, region, q1, q2, None, None, renewal, expansion))
    return rows


def write_formulas(ws, row_num, broken):
    if broken:
        ws[f'E{row_num}'] = f'=C{row_num}-D{row_num}'
        ws[f'F{row_num}'] = f'=E{row_num}/C{row_num}'
    else:
        ws[f'E{row_num}'] = f'=C{row_num}+D{row_num}'
        ws[f'F{row_num}'] = f'=IF(C{row_num}=0,0,E{row_num}/C{row_num}-1)'


def apply_full_styles(ws, row_count):
    apply_template_layout(ws, row_count)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center')
    ws.row_dimensions[1].height = 24
    for row_num in range(2, row_count + 2):
        for col in ('C', 'D', 'E'):
            ws[f'{col}{row_num}'].number_format = '$#,##0'
            ws[f'{col}{row_num}'].font = DATA_FONT
        for col in ('F', 'G', 'H'):
            ws[f'{col}{row_num}'].number_format = '0.0%'
            ws[f'{col}{row_num}'].font = DATA_FONT
        for col in ('E', 'F'):
            ws[f'{col}{row_num}'].fill = TARGET_FILL
            ws[f'{col}{row_num}'].alignment = Alignment(horizontal='center')
            ws[f'{col}{row_num}'].border = TARGET_BORDER


def apply_template_layout(ws, row_count):
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f'A1:H{row_count + 1}'
    for column_name, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[column_name].width = width


def apply_local_emphasis(ws, row_count):
    for row_num in range(2, row_count + 2):
        for col in ('G', 'H'):
            ws[f'{col}{row_num}'].fill = LOCAL_FILL
            ws[f'{col}{row_num}'].alignment = Alignment(horizontal='right')
        ws[f'G{row_num}'].font = Font(color='FF1F1F1F', bold=True)
        ws[f'H{row_num}'].font = Font(color='FF1F1F1F', italic=True)


def corrupt_target_styles(ws, row_count):
    ws.freeze_panes = 'A1'
    ws.auto_filter.ref = None
    ws.row_dimensions[1].height = 18
    ws.column_dimensions['E'].width = 10.0
    ws.column_dimensions['F'].width = 9.0
    for row_num in range(2, row_count + 2):
        for col in ('E', 'F'):
            ws[f'{col}{row_num}'].number_format = 'General'
            ws[f'{col}{row_num}'].fill = PatternFill()
            ws[f'{col}{row_num}'].font = Font(color='FFFF0000')
            ws[f'{col}{row_num}'].alignment = Alignment(horizontal='left')
            ws[f'{col}{row_num}'].border = Border()


def build_workbook(output_file):
    wb = Workbook()
    reference = wb.active
    reference.title = 'Reference'
    target = wb.create_sheet('BoardKPI')
    reference.append(['Team', 'Region', 'Q1 Revenue', 'Q2 Revenue', 'H1 Revenue', 'H1 Growth %', 'Renewal Rate', 'Expansion Rate'])
    target.append(['Team', 'Region', 'Q1 Revenue', 'Q2 Revenue', 'H1 Revenue', 'H1 Growth %', 'Renewal Rate', 'Expansion Rate'])
    for row in build_rows(GUIDE_ROW_COUNT, 1):
        reference.append(list(row))
    for row in build_rows(DATA_ROW_COUNT, GUIDE_ROW_COUNT + 1):
        target.append(list(row))
    for row_num in range(2, GUIDE_ROW_COUNT + 2):
        write_formulas(reference, row_num, broken=False)
    for row_num in range(2, DATA_ROW_COUNT + 2):
        write_formulas(target, row_num, broken=True)
    apply_full_styles(reference, GUIDE_ROW_COUNT)
    apply_full_styles(target, DATA_ROW_COUNT)
    apply_local_emphasis(target, DATA_ROW_COUNT)
    corrupt_target_styles(target, DATA_ROW_COUNT)
    wb.save(output_file)
    wb.close()


def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    output_file = os.path.join(args.output_path, 'board_kpi_style.xlsx')
    build_workbook(output_file)
    with open(os.path.join(args.output_path, 'task_metadata.json'), 'w', encoding='utf-8') as f:
        json.dump({'task_id': 'style_preserve/task_001_restore_board_kpi_styles_after_formula_fix', 'target_file': 'board_kpi_style.xlsx'}, f, ensure_ascii=False, indent=2)
    print(os.path.abspath(args.output_path))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Create workspace for style preserve task')
    parser.add_argument('output_path', help='Directory to create workspace in')
    main(parser.parse_args())
