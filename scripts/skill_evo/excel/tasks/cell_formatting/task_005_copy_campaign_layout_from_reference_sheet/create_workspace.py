import argparse
import json
import os
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


DATA_ROW_COUNT = 200
GUIDE_ROW_COUNT = 20
STATES = ['Ready', 'Blocked', 'Pending Legal', 'Review Needed']
CHANNELS = ['Search', 'Social', 'Affiliate', 'Email']
HEADER_FILL = PatternFill('solid', fgColor='134F5C')
STATE_FILL = PatternFill('solid', fgColor='FCE5CD')
HEADER_FONT = Font(bold=True, color='FFFFFFFF')
DATA_FONT = Font(color='FF000000')


def build_rows(count, start_index=1):
    base_date = date(2026, 4, 1)
    rows = []
    for i in range(start_index, start_index + count):
        rows.append([
            f'{CHANNELS[(i - 1) % len(CHANNELS)]}-{i:04d}',
            base_date + timedelta(days=i - 1),
            round(2500 + i * 81.25, 2),
            round(0.012 + (i % 9) * 0.0045, 4),
            STATES[(i - 1) % len(STATES)],
        ])
    return rows


def apply_styles(ws, data_row_count):
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center')
    for row_num in range(2, data_row_count + 2):
        ws[f'B{row_num}'].number_format = 'yyyy-mm-dd'
        ws[f'B{row_num}'].font = DATA_FONT
        ws[f'C{row_num}'].number_format = '$#,##0.00'
        ws[f'C{row_num}'].font = DATA_FONT
        ws[f'D{row_num}'].number_format = '0.0%'
        ws[f'D{row_num}'].font = DATA_FONT
        ws[f'E{row_num}'].fill = STATE_FILL
        ws[f'E{row_num}'].alignment = Alignment(horizontal='center')


def build_workbook(output_file):
    wb = Workbook()
    guide_ws = wb.active
    guide_ws.title = 'Reference Layout'
    target_ws = wb.create_sheet('Launch Queue')
    headers = ['Campaign ID', 'Launch Date', 'Planned Spend', 'CTR Goal', 'Approval State']
    guide_ws.append(headers)
    target_ws.append(headers)
    for row in build_rows(GUIDE_ROW_COUNT, 1):
        guide_ws.append(row)
    for row in build_rows(DATA_ROW_COUNT, GUIDE_ROW_COUNT + 1):
        target_ws.append(row)
    apply_styles(guide_ws, GUIDE_ROW_COUNT)
    wb.save(output_file)
    wb.close()


def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    output_file = os.path.join(args.output_path, 'campaign_launch_queue.xlsx')
    build_workbook(output_file)
    with open(os.path.join(args.output_path, 'task_metadata.json'), 'w', encoding='utf-8') as f:
        json.dump(
            {
                'task_id': 'cell_formatting/task_005_copy_campaign_layout_from_reference_sheet',
                'target_file': 'campaign_launch_queue.xlsx',
            },
            f,
            ensure_ascii=False,
            indent=2,
        )
    print(os.path.abspath(args.output_path))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Create workspace for cell formatting task')
    parser.add_argument('output_path', help='Directory to create workspace in')
    main(parser.parse_args())
