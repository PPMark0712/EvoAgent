import argparse
import json
import os
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


TARGET_ROW_COUNT = 200
REFERENCE_ROW_COUNT = 20
SEVERITIES = ['Low', 'Medium', 'High', 'Critical']
OWNERS = ['CS Ops', 'Finance', 'Sales Ops', 'RevOps']
HEADER_FILL = PatternFill('solid', fgColor='7F6000')
SEVERITY_FILL = PatternFill('solid', fgColor='F4CCCC')
HEADER_FONT = Font(bold=True, color='FFFFFFFF')
DATA_FONT = Font(color='FF000000')


def build_rows(count, start_index=1):
    base_date = date(2026, 2, 1)
    rows = []
    for i in range(start_index, start_index + count):
        rows.append([
            f'EXC-{i:04d}',
            base_date + timedelta(days=i - 1),
            round(1200 + i * 24.5, 2),
            round(0.01 + (i % 8) * 0.006, 4),
            SEVERITIES[(i - 1) % len(SEVERITIES)],
        ])
    return rows


def apply_styles(ws, header_row, data_start_row, data_row_count, start_col):
    for col_idx in range(start_col, start_col + 5):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center')
    for row_num in range(data_start_row, data_start_row + data_row_count):
        ws.cell(row=row_num, column=start_col + 1).number_format = 'yyyy-mm-dd'
        ws.cell(row=row_num, column=start_col + 1).font = DATA_FONT
        ws.cell(row=row_num, column=start_col + 2).number_format = '$#,##0.00'
        ws.cell(row=row_num, column=start_col + 2).font = DATA_FONT
        ws.cell(row=row_num, column=start_col + 3).number_format = '0.0%'
        ws.cell(row=row_num, column=start_col + 3).font = DATA_FONT
        ws.cell(row=row_num, column=start_col + 4).fill = SEVERITY_FILL
        ws.cell(row=row_num, column=start_col + 4).alignment = Alignment(horizontal='center')


def build_workbook(output_file):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Exceptions'
    target_headers = ['Exception ID', 'Review Date', 'Exposure Amount', 'Recovery Rate', 'Severity']
    ref_headers = list(target_headers)
    ws.append(target_headers)
    for row in build_rows(TARGET_ROW_COUNT, 1):
        ws.append(row)
    ws['G1'] = ref_headers[0]
    ws['H1'] = ref_headers[1]
    ws['I1'] = ref_headers[2]
    ws['J1'] = ref_headers[3]
    ws['K1'] = ref_headers[4]
    for offset, row in enumerate(build_rows(REFERENCE_ROW_COUNT, TARGET_ROW_COUNT + 1), start=2):
        for col_offset, value in enumerate(row, start=7):
            ws.cell(row=offset, column=col_offset, value=value)
    ws['G23'] = 'Use the styled block above as the formatting reference.'
    ws['G24'] = 'Do not change any values.'
    apply_styles(ws, 1, 2, REFERENCE_ROW_COUNT, 7)
    wb.save(output_file)
    wb.close()


def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    output_file = os.path.join(args.output_path, 'renewal_exception_log.xlsx')
    build_workbook(output_file)
    with open(os.path.join(args.output_path, 'task_metadata.json'), 'w', encoding='utf-8') as f:
        json.dump(
            {
                'task_id': 'cell_formatting/task_007_extend_exception_log_style_within_sheet',
                'target_file': 'renewal_exception_log.xlsx',
            },
            f,
            ensure_ascii=False,
            indent=2,
        )
    print(os.path.abspath(args.output_path))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Create workspace for cell formatting task')
    parser.add_argument('output_path', help='Directory to create workspace in')
    main(parser.parse_args())
