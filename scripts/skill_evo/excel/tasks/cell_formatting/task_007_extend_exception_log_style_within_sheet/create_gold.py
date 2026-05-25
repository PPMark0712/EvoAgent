import argparse
import os

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


TARGET_ROW_COUNT = 200
REFERENCE_ROW_COUNT = 20
HEADER_FILL = PatternFill('solid', fgColor='7F6000')
SEVERITY_FILL = PatternFill('solid', fgColor='F4CCCC')
HEADER_FONT = Font(bold=True, color='FFFFFFFF')
DATA_FONT = Font(color='FF000000')


def build_workbook(output_file):
    from create_workspace import build_workbook as create_input_workbook

    create_input_workbook(output_file)


def apply_styles(ws, header_row, data_start_row, data_row_count, start_col):
    for col_idx in range(start_col, start_col + 5):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center')
    for row_num in range(data_start_row, data_start_row + data_row_count):
        ws.cell(row=row_num, column=start_col + 1).number_format = 'yyyy-mm-dd'
        ws.cell(row=row_num, column=start_col + 1).font = DATA_FONT
        ws.cell(row=row_num, column=start_col + 2).number_format = '$#,##0.00'
        ws.cell(row=row_num, column=start_col + 2).font = DATA_FONT
        ws.cell(row=row_num, column=start_col + 3).number_format = '0.0%'
        ws.cell(row=row_num, column=start_col + 3).font = DATA_FONT
        ws.cell(row=row_num, column=start_col + 4).fill = SEVERITY_FILL
        ws.cell(row=row_num, column=start_col + 4).alignment = Alignment(horizontal='center')


def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    target = os.path.join(args.output_path, 'renewal_exception_log.xlsx')
    if not os.path.isfile(target):
        build_workbook(target)
    wb = load_workbook(target)
    ws = wb['Exceptions']
    apply_styles(ws, 1, 2, TARGET_ROW_COUNT, 1)
    apply_styles(ws, 1, 2, REFERENCE_ROW_COUNT, 7)
    wb.save(target)
    wb.close()
    print(os.path.abspath(target))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Create gold workbook for cell formatting task')
    parser.add_argument('output_path', help='Directory to place gold workbook in')
    main(parser.parse_args())
