import argparse
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


TARGET_ROW_COUNT = 200
REFERENCE_ROW_COUNT = 20
HEADER_FILL = PatternFill('solid', fgColor='5B9BD5')
READINESS_FILL = PatternFill('solid', fgColor='FFF2CC')
HEADER_FONT = Font(bold=True, color='FFFFFFFF')
DATA_FONT = Font(color='FF000000')


def build_rows(count, start_index=1):
    rows = []
    for i in range(start_index, start_index + count):
        rows.append([
            f'Team-{i:03d}',
            12 + (i % 9),
            round(0.64 + (i % 7) * 0.03, 4),
            18000 + i * 220,
            ['Stable', 'Ramp', 'Constraint', 'Review'][(i - 1) % 4],
        ])
    return rows


def apply_styles(ws, header_row, data_start_row, data_row_count):
    for cell in ws[header_row]:
        if cell.column > 5:
            continue
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center')
    for row_num in range(data_start_row, data_start_row + data_row_count):
        ws[f'B{row_num}'].number_format = '0'
        ws[f'B{row_num}'].font = DATA_FONT
        ws[f'C{row_num}'].number_format = '0.0%'
        ws[f'C{row_num}'].font = DATA_FONT
        ws[f'D{row_num}'].number_format = '$#,##0'
        ws[f'D{row_num}'].font = DATA_FONT
        ws[f'E{row_num}'].fill = READINESS_FILL
        ws[f'E{row_num}'].alignment = Alignment(horizontal='center')


def build_workbook(output_file):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Capacity Plan'
    headers = ['Team', 'Open Seats', 'Utilization Target', 'Hiring Budget', 'Readiness']
    ws.append(headers)
    for row in build_rows(REFERENCE_ROW_COUNT, 1):
        ws.append(row)
    ws['A22'] = 'Reference block ends above. Keep it unchanged.'
    ws['A23'] = 'Format the new block below by matching the style pattern.'
    ws.append(headers)
    for row in build_rows(TARGET_ROW_COUNT, REFERENCE_ROW_COUNT + 1):
        ws.append(row)
    apply_styles(ws, 1, 2, REFERENCE_ROW_COUNT)
    wb.save(output_file)
    wb.close()


def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    output_file = os.path.join(args.output_path, 'quarterly_capacity_plan.xlsx')
    build_workbook(output_file)
    with open(os.path.join(args.output_path, 'task_metadata.json'), 'w', encoding='utf-8') as f:
        json.dump(
            {
                'task_id': 'cell_formatting/task_006_extend_capacity_plan_style_within_sheet',
                'target_file': 'quarterly_capacity_plan.xlsx',
            },
            f,
            ensure_ascii=False,
            indent=2,
        )
    print(os.path.abspath(args.output_path))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Create workspace for cell formatting task')
    parser.add_argument('output_path', help='Directory to create workspace in')
    main(parser.parse_args())
