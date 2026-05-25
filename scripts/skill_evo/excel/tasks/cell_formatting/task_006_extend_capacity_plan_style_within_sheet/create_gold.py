import argparse
import os

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


REFERENCE_ROW_COUNT = 20
TARGET_ROW_COUNT = 200
REFERENCE_HEADER_ROW = 1
TARGET_HEADER_ROW = 24
TARGET_DATA_START_ROW = 25
HEADER_FILL = PatternFill('solid', fgColor='5B9BD5')
READINESS_FILL = PatternFill('solid', fgColor='FFF2CC')
HEADER_FONT = Font(bold=True, color='FFFFFFFF')
DATA_FONT = Font(color='FF000000')


def build_workbook(output_file):
    from create_workspace import build_workbook as create_input_workbook

    create_input_workbook(output_file)


def apply_styles(ws, header_row, data_start_row, data_row_count):
    for cell in ws[header_row]:
        if cell.column > 5:
            continue
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center')
    for row_num in range(data_start_row, data_start_row + data_row_count):
        ws[f'B{row_num}'].number_format = '0'
        ws[f'B{row_num}'].font = DATA_FONT
        ws[f'C{row_num}'].number_format = '0.0%'
        ws[f'C{row_num}'].font = DATA_FONT
        ws[f'D{row_num}'].number_format = '$#,##0'
        ws[f'D{row_num}'].font = DATA_FONT
        ws[f'E{row_num}'].fill = READINESS_FILL
        ws[f'E{row_num}'].alignment = Alignment(horizontal='center')


def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    target = os.path.join(args.output_path, 'quarterly_capacity_plan.xlsx')
    if not os.path.isfile(target):
        build_workbook(target)
    wb = load_workbook(target)
    ws = wb['Capacity Plan']
    apply_styles(ws, REFERENCE_HEADER_ROW, 2, REFERENCE_ROW_COUNT)
    apply_styles(ws, TARGET_HEADER_ROW, TARGET_DATA_START_ROW, TARGET_ROW_COUNT)
    wb.save(target)
    wb.close()
    print(os.path.abspath(target))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Create gold workbook for cell formatting task')
    parser.add_argument('output_path', help='Directory to place gold workbook in')
    main(parser.parse_args())
