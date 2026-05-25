import argparse
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


DATA_ROW_COUNT = 200
GUIDE_ROW_COUNT = 20
STATUSES = ['On Track', 'Needs Review', 'Stretch', 'Watchlist']
REGIONS = ['North America', 'EMEA', 'APAC', 'LATAM']
HEADER_FILL = PatternFill('solid', fgColor='203864')
STATUS_FILL = PatternFill('solid', fgColor='D9EAD3')
HEADER_FONT = Font(bold=True, color='FFFFFFFF')
DATA_FONT = Font(color='FF000000')


def build_rows(count, start_index=1):
    rows = []
    for i in range(start_index, start_index + count):
        rows.append([
            REGIONS[(i - 1) % len(REGIONS)],
            85000 + i * 1450,
            round(0.04 + (i % 8) * 0.011, 4),
            round(0.72 + (i % 6) * 0.03, 4),
            STATUSES[(i - 1) % len(STATUSES)],
        ])
    return rows


def apply_styles(ws, data_row_count):
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center')
    for row_num in range(2, data_row_count + 2):
        ws[f'B{row_num}'].number_format = '$#,##0'
        ws[f'B{row_num}'].font = DATA_FONT
        ws[f'C{row_num}'].number_format = '0.0%'
        ws[f'C{row_num}'].font = DATA_FONT
        ws[f'D{row_num}'].number_format = '0.0%'
        ws[f'D{row_num}'].font = DATA_FONT
        ws[f'E{row_num}'].fill = STATUS_FILL
        ws[f'E{row_num}'].alignment = Alignment(horizontal='center')


def build_workbook(output_file):
    wb = Workbook()
    guide_ws = wb.active
    guide_ws.title = 'Style Guide'
    target_ws = wb.create_sheet('Targets')
    headers = ['Region', 'ARR Plan', 'Growth Target', 'Renewal Target', 'Status']
    guide_ws.append(headers)
    target_ws.append(headers)
    for row in build_rows(GUIDE_ROW_COUNT, 1):
        guide_ws.append(row)
    for row in build_rows(DATA_ROW_COUNT, GUIDE_ROW_COUNT + 1):
        target_ws.append(row)
    apply_styles(guide_ws, GUIDE_ROW_COUNT)
    wb.save(output_file)
    wb.close()


def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    output_file = os.path.join(args.output_path, 'board_pack_targets.xlsx')
    build_workbook(output_file)
    with open(os.path.join(args.output_path, 'task_metadata.json'), 'w', encoding='utf-8') as f:
        json.dump(
            {
                'task_id': 'cell_formatting/task_004_copy_board_pack_style_from_reference_sheet',
                'target_file': 'board_pack_targets.xlsx',
            },
            f,
            ensure_ascii=False,
            indent=2,
        )
    print(os.path.abspath(args.output_path))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Create workspace for cell formatting task')
    parser.add_argument('output_path', help='Directory to create workspace in')
    main(parser.parse_args())
