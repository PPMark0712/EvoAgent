import argparse
import os

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


DATA_ROW_COUNT = 200
GUIDE_ROW_COUNT = 20
HEADER_FILL = PatternFill('solid', fgColor='203864')
STATUS_FILL = PatternFill('solid', fgColor='D9EAD3')
HEADER_FONT = Font(bold=True, color='FFFFFFFF')
DATA_FONT = Font(color='FF000000')


def build_workbook(output_file):
    from create_workspace import build_workbook as create_input_workbook

    create_input_workbook(output_file)


def apply_styles(ws, data_row_count):
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center')
    for row_num in range(2, data_row_count + 2):
        ws[f'B{row_num}'].number_format = '$#,##0'
        ws[f'B{row_num}'].font = DATA_FONT
        ws[f'C{row_num}'].number_format = '0.0%'
        ws[f'C{row_num}'].font = DATA_FONT
        ws[f'D{row_num}'].number_format = '0.0%'
        ws[f'D{row_num}'].font = DATA_FONT
        ws[f'E{row_num}'].fill = STATUS_FILL
        ws[f'E{row_num}'].alignment = Alignment(horizontal='center')


def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    target = os.path.join(args.output_path, 'board_pack_targets.xlsx')
    if not os.path.isfile(target):
        build_workbook(target)
    wb = load_workbook(target)
    apply_styles(wb['Style Guide'], GUIDE_ROW_COUNT)
    apply_styles(wb['Targets'], DATA_ROW_COUNT)
    wb.save(target)
    wb.close()
    print(os.path.abspath(target))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Create gold workbook for cell formatting task')
    parser.add_argument('output_path', help='Directory to place gold workbook in')
    main(parser.parse_args())
