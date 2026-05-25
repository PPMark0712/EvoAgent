import argparse
import os

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


DATA_ROW_COUNT = 200
HEADER_FILL = PatternFill("solid", fgColor="38761D")
STATUS_FILL = PatternFill("solid", fgColor="F4CCCC")
HEADER_FONT = Font(bold=True, color="FFFFFFFF")
DATA_FONT = Font(color="FF000000")


def build_workbook(output_file):
    from create_workspace import build_workbook as create_input_workbook

    create_input_workbook(output_file)



def apply_styles(ws):
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    for row_num in range(2, DATA_ROW_COUNT + 2):
        ws[f"B{row_num}"].number_format = "yyyy-mm-dd"
        ws[f"B{row_num}"].font = DATA_FONT
        ws[f"C{row_num}"].number_format = "$#,##0.00"
        ws[f"C{row_num}"].font = DATA_FONT
        ws[f"D{row_num}"].number_format = "0.0%"
        ws[f"D{row_num}"].font = DATA_FONT
        ws[f"E{row_num}"].fill = STATUS_FILL
        ws[f"E{row_num}"].alignment = Alignment(horizontal="center")



def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    target = os.path.join(args.output_path, "expense_audit.xlsx")
    if not os.path.isfile(target):
        build_workbook(target)
    wb = load_workbook(target)
    apply_styles(wb["Expenses"])
    wb.save(target)
    wb.close()
    print(os.path.abspath(target))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Create gold workbook for cell formatting task")
    parser.add_argument("output_path", help="Directory to place gold workbook in")
    main(parser.parse_args())
