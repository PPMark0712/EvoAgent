import argparse
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Font


DATA_ROW_COUNT = 200


def canonical_rows():
    rows = []
    for i in range(1, DATA_ROW_COUNT + 1):
        growth_rate = round(0.04 + (i % 12) * 0.005, 4)
        marketing_budget = 120 + i * 3
        headcount = 8 + (i % 15)
        revenue_target = 900 + i * 22
        rows.append([
            f"SCN-{i:03d}",
            growth_rate,
            marketing_budget,
            headcount,
            revenue_target,
        ])
    return rows


def add_summary_row(ws):
    summary_row = DATA_ROW_COUNT + 2
    ws[f"A{summary_row}"] = "Portfolio Summary"
    ws[f"B{summary_row}"] = f"=AVERAGE(B2:B{DATA_ROW_COUNT + 1})"
    ws[f"C{summary_row}"] = f"=SUM(C2:C{DATA_ROW_COUNT + 1})"
    ws[f"D{summary_row}"] = f"=AVERAGE(D2:D{DATA_ROW_COUNT + 1})"
    ws[f"E{summary_row}"] = f"=MAX(E2:E{DATA_ROW_COUNT + 1})"


def build_workbook(output_file):
    wb = Workbook()
    ws = wb.active
    ws.title = "Inputs"
    headers = ["Scenario ID", "Growth Rate", "Marketing Budget", "Headcount", "Revenue Target"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in canonical_rows():
        ws.append(row)
    notes = wb.create_sheet("Notes")
    notes["A1"] = "Planning Notes"
    notes["A2"] = "Keep the 200 scenario rows unchanged and append one summary row."
    notes["A3"] = "Summary metrics should be created with Excel formulas, not hardcoded numbers."
    notes["A1"].font = Font(bold=True)
    wb.save(output_file)
    wb.close()


def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    output_file = os.path.join(args.output_path, "budget_update.xlsx")
    build_workbook(output_file)
    with open(os.path.join(args.output_path, "task_metadata.json"), "w", encoding="utf-8") as f:
        json.dump(
            {
                "task_id": "direct_cell_update/task_001_adjust_assumptions",
                "target_file": "budget_update.xlsx",
            },
            f,
            ensure_ascii=False,
            indent=2,
        )
    print(os.path.abspath(args.output_path))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Create workspace for direct cell update task")
    parser.add_argument("output_path", help="Directory to create workspace in")
    main(parser.parse_args())
