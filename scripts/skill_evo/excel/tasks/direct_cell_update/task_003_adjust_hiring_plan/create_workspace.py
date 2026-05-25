import argparse
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Font


DATA_ROW_COUNT = 200


def canonical_rows():
    rows = []
    for i in range(1, DATA_ROW_COUNT + 1):
        quarterly_hires = 2 + (i % 8)
        attrition_rate = round(0.01 + (i % 7) * 0.006, 4)
        training_budget = 18 + (i % 11) * 4
        ending_headcount = 35 + i
        rows.append([
            f"TEAM-{i:03d}",
            quarterly_hires,
            attrition_rate,
            training_budget,
            ending_headcount,
        ])
    return rows


def add_summary_row(ws):
    summary_row = DATA_ROW_COUNT + 2
    ws[f"A{summary_row}"] = "Hiring Summary"
    ws[f"B{summary_row}"] = f"=SUM(B2:B{DATA_ROW_COUNT + 1})"
    ws[f"C{summary_row}"] = f"=AVERAGE(C2:C{DATA_ROW_COUNT + 1})"
    ws[f"D{summary_row}"] = f"=SUM(D2:D{DATA_ROW_COUNT + 1})"
    ws[f"E{summary_row}"] = f"=AVERAGE(E2:E{DATA_ROW_COUNT + 1})"


def build_workbook(output_file):
    wb = Workbook()
    ws = wb.active
    ws.title = "Hiring"
    headers = ["Team ID", "Quarterly Hires", "Attrition Rate", "Training Budget", "Ending Headcount"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in canonical_rows():
        ws.append(row)
    notes = wb.create_sheet("Notes")
    notes["A1"] = "Workforce Plan"
    notes["A2"] = "Append a single summary row after the 200 team records."
    notes["A3"] = "The statistical outputs must be formulas, not copied numbers."
    notes["A1"].font = Font(bold=True)
    wb.save(output_file)
    wb.close()


def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    output_file = os.path.join(args.output_path, "workforce_plan.xlsx")
    build_workbook(output_file)
    with open(os.path.join(args.output_path, "task_metadata.json"), "w", encoding="utf-8") as f:
        json.dump(
            {
                "task_id": "direct_cell_update/task_003_adjust_hiring_plan",
                "target_file": "workforce_plan.xlsx",
            },
            f,
            ensure_ascii=False,
            indent=2,
        )
    print(os.path.abspath(args.output_path))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Create workspace for direct cell update task")
    parser.add_argument("output_path", help="Directory to create workspace in")
    main(parser.parse_args())
