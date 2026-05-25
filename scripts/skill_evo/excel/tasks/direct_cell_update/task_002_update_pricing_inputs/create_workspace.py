import argparse
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Font


DATA_ROW_COUNT = 200


def canonical_rows():
    rows = []
    for i in range(1, DATA_ROW_COUNT + 1):
        base_price = 79 + (i % 14) * 3
        discount_rate = round(0.02 + (i % 6) * 0.01, 4)
        renewal_rate = round(0.72 + (i % 10) * 0.015, 4)
        support_cost = 12 + (i % 9) * 2
        rows.append([
            f"PLAN-{i:03d}",
            base_price,
            discount_rate,
            renewal_rate,
            support_cost,
        ])
    return rows


def add_summary_row(ws):
    summary_row = DATA_ROW_COUNT + 2
    ws[f"A{summary_row}"] = "Pricing Summary"
    ws[f"B{summary_row}"] = f"=AVERAGE(B2:B{DATA_ROW_COUNT + 1})"
    ws[f"C{summary_row}"] = f"=AVERAGE(C2:C{DATA_ROW_COUNT + 1})"
    ws[f"D{summary_row}"] = f"=AVERAGE(D2:D{DATA_ROW_COUNT + 1})"
    ws[f"E{summary_row}"] = f"=SUM(E2:E{DATA_ROW_COUNT + 1})"


def build_workbook(output_file):
    wb = Workbook()
    ws = wb.active
    ws.title = "Pricing"
    headers = ["Plan ID", "Base Price", "Discount Rate", "Renewal Rate", "Support Cost"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in canonical_rows():
        ws.append(row)
    notes = wb.create_sheet("Notes")
    notes["A1"] = "Pricing Review"
    notes["A2"] = "Add one summary row after the last pricing record."
    notes["A3"] = "Use formulas for the statistics instead of typing fixed numbers."
    notes["A1"].font = Font(bold=True)
    wb.save(output_file)
    wb.close()


def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    output_file = os.path.join(args.output_path, "pricing_model.xlsx")
    build_workbook(output_file)
    with open(os.path.join(args.output_path, "task_metadata.json"), "w", encoding="utf-8") as f:
        json.dump(
            {
                "task_id": "direct_cell_update/task_002_update_pricing_inputs",
                "target_file": "pricing_model.xlsx",
            },
            f,
            ensure_ascii=False,
            indent=2,
        )
    print(os.path.abspath(args.output_path))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Create workspace for direct cell update task")
    parser.add_argument("output_path", help="Directory to create workspace in")
    main(parser.parse_args())
