import argparse
import os

from openpyxl import load_workbook

from create_workspace import ROW_COUNT, build_workbook


def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    target = os.path.join(args.output_path, 'quarterly_plan.xlsx')
    if not os.path.isfile(target):
        build_workbook(target)
    wb = load_workbook(target)
    ws = wb['Plan']
    for row_num in range(2, ROW_COUNT + 2):
        ws[f'E{row_num}'] = f'=Drivers!E{row_num}*(1+Drivers!G{row_num})'
        ws[f'F{row_num}'] = f'=Drivers!F{row_num}*(1+Drivers!I{row_num})'
        ws[f'G{row_num}'] = f'=E{row_num}*F{row_num}'
        ws[f'H{row_num}'] = f'=E{row_num}*(1+Drivers!H{row_num})'
        ws[f'I{row_num}'] = f'=F{row_num}*(1+Drivers!J{row_num})'
        ws[f'J{row_num}'] = f'=H{row_num}*I{row_num}'
        ws[f'K{row_num}'] = f'=G{row_num}+J{row_num}'
        ws[f'L{row_num}'] = f'=IF(D{row_num}=0,0,(K{row_num}-D{row_num})/D{row_num})'
    wb.save(target)
    wb.close()
    print(os.path.abspath(target))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Create gold workbook for formula repair task')
    parser.add_argument('output_path', help='Directory to place gold workbook in')
    main(parser.parse_args())
