import argparse
import os

from openpyxl import load_workbook
from openpyxl.styles import Font


ROW_COUNT = 220
TOTAL_ROW = ROW_COUNT + 2


def build_workbook(output_file):
    from create_workspace import build_workbook as create_input_workbook

    create_input_workbook(output_file)


def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    target = os.path.join(args.output_path, "subscription_report.xlsx")
    if not os.path.isfile(target):
        build_workbook(target)
    wb = load_workbook(target)
    ws = wb["Summary"]
    for row_num in range(2, ROW_COUNT + 2):
        ws[f"C{row_num}"] = f"=Subscribers!C{row_num}+Subscribers!D{row_num}+Subscribers!E{row_num}-Subscribers!F{row_num}"
        ws[f"D{row_num}"] = f"=Subscribers!D{row_num}+Subscribers!E{row_num}"
        ws[f"E{row_num}"] = f"=Subscribers!F{row_num}"
        ws[f"F{row_num}"] = f"=C{row_num}*Subscribers!G{row_num}+Subscribers!J{row_num}"
        ws[f"G{row_num}"] = f"=C{row_num}*Subscribers!H{row_num}"
        ws[f"H{row_num}"] = f"=F{row_num}*Subscribers!I{row_num}"
        ws[f"I{row_num}"] = f"=F{row_num}-G{row_num}-H{row_num}"
        ws[f"J{row_num}"] = f"=IF(Subscribers!C{row_num}=0,0,C{row_num}/Subscribers!C{row_num})"
    for coord in [f"A{TOTAL_ROW}", f"B{TOTAL_ROW}", f"C{TOTAL_ROW}", f"D{TOTAL_ROW}", f"E{TOTAL_ROW}", f"F{TOTAL_ROW}", f"G{TOTAL_ROW}", f"H{TOTAL_ROW}", f"I{TOTAL_ROW}", f"J{TOTAL_ROW}"]:
        ws[coord].font = Font(bold=True)
    ws[f"B{TOTAL_ROW}"] = "All Plans"
    ws[f"C{TOTAL_ROW}"] = f"=SUM(C2:C{ROW_COUNT + 1})"
    ws[f"D{TOTAL_ROW}"] = f"=SUM(D2:D{ROW_COUNT + 1})"
    ws[f"E{TOTAL_ROW}"] = f"=SUM(E2:E{ROW_COUNT + 1})"
    ws[f"F{TOTAL_ROW}"] = f"=SUM(F2:F{ROW_COUNT + 1})"
    ws[f"G{TOTAL_ROW}"] = f"=SUM(G2:G{ROW_COUNT + 1})"
    ws[f"H{TOTAL_ROW}"] = f"=SUM(H2:H{ROW_COUNT + 1})"
    ws[f"I{TOTAL_ROW}"] = f"=SUM(I2:I{ROW_COUNT + 1})"
    ws[f"J{TOTAL_ROW}"] = f"=AVERAGE(J2:J{ROW_COUNT + 1})"
    wb.save(target)
    wb.close()
    print(os.path.abspath(target))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Create gold workbook for formula repair task")
    parser.add_argument("output_path", help="Directory to place gold workbook in")
    main(parser.parse_args())
