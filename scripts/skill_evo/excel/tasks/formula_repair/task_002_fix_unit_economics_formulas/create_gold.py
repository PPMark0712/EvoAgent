import argparse
import os

from openpyxl import load_workbook
from openpyxl.styles import Font


ROW_COUNT = 220


def build_workbook(output_file):
    from create_workspace import build_workbook as create_input_workbook

    create_input_workbook(output_file)


def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    target = os.path.join(args.output_path, "unit_economics.xlsx")
    if not os.path.isfile(target):
        build_workbook(target)
    wb = load_workbook(target)
    ws = wb["Model"]
    ws["K1"] = "Contribution Margin %"
    ws["K1"].font = Font(bold=True)
    for row_num in range(2, ROW_COUNT + 2):
        ws[f"C{row_num}"] = f"=Raw!D{row_num}*Raw!E{row_num}"
        ws[f"D{row_num}"] = f"=Raw!C{row_num}*Raw!F{row_num}"
        ws[f"E{row_num}"] = f"=Raw!D{row_num}*Raw!G{row_num}"
        ws[f"F{row_num}"] = f"=C{row_num}*Raw!H{row_num}"
        ws[f"G{row_num}"] = f"=C{row_num}*Raw!I{row_num}"
        ws[f"H{row_num}"] = f"=C{row_num}-F{row_num}-G{row_num}"
        ws[f"I{row_num}"] = f"=H{row_num}-D{row_num}-E{row_num}"
        ws[f"J{row_num}"] = f"=IF(D{row_num}=0,0,I{row_num}/D{row_num})"
        ws[f"K{row_num}"] = f"=IF(H{row_num}=0,0,I{row_num}/H{row_num})"
    wb.save(target)
    wb.close()
    print(os.path.abspath(target))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Create gold workbook for formula repair task")
    parser.add_argument("output_path", help="Directory to place gold workbook in")
    main(parser.parse_args())
