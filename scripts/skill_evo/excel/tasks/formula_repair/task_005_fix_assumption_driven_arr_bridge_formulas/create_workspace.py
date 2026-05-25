import argparse
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Font


ROW_COUNT = 220
SEGMENTS = ['SMB', 'Mid-Market', 'Enterprise', 'Strategic']
ASSUMPTION_VALUES = {
    'Upsell Rate': 0.08,
    'Risk Discount': 0.06,
    'Cost per Support Hour': 12,
    'Retention Bonus': 180,
}


def canonical_rows():
    rows = []
    for idx in range(ROW_COUNT):
        current_arr = 4200 + (idx % 24) * 210
        expansion_arr = 260 + (idx % 9) * 35
        churned_arr = 180 + (idx % 7) * 20
        renewal_probability = 0.72 + (idx % 5) * 0.04
        support_hours = 8 + (idx % 6)
        fee_rate = 0.018 + (idx % 4) * 0.004
        growth_uplift = 0.01 + (idx % 5) * 0.01
        onboarding_credit = 90 + (idx % 4) * 25
        rows.append((
            f'ACCT-{idx + 1:03d}',
            SEGMENTS[idx % len(SEGMENTS)],
            current_arr,
            expansion_arr,
            churned_arr,
            renewal_probability,
            support_hours,
            fee_rate,
            growth_uplift,
            onboarding_credit,
        ))
    return rows


def build_workbook(output_file):
    wb = Workbook()
    accounts = wb.active
    accounts.title = 'Accounts'
    account_headers = [
        'Account ID',
        'Segment',
        'Current ARR',
        'Expansion ARR',
        'Churned ARR',
        'Renewal Probability',
        'Support Hours',
        'Fee Rate',
        'Growth Uplift',
        'Onboarding Credit',
    ]
    accounts.append(account_headers)
    for cell in accounts[1]:
        cell.font = Font(bold=True)
    for row in canonical_rows():
        accounts.append(list(row))

    assumptions = wb.create_sheet('Assumptions')
    assumptions['A1'] = 'Driver'
    assumptions['B1'] = 'Value'
    assumptions['A1'].font = Font(bold=True)
    assumptions['B1'].font = Font(bold=True)
    for idx, (label, value) in enumerate(ASSUMPTION_VALUES.items(), start=2):
        assumptions[f'A{idx}'] = label
        assumptions[f'B{idx}'] = value

    bridge = wb.create_sheet('Bridge')
    bridge_headers = [
        'Account ID',
        'Segment',
        'Gross ARR',
        'Renewal ARR',
        'Upsell ARR',
        'Risk Adjusted ARR',
        'Support Cost',
        'Payment Fees',
        'Net ARR',
        'Net ARR Margin %',
    ]
    bridge.append(bridge_headers)
    for cell in bridge[1]:
        cell.font = Font(bold=True)
    for row_num in range(2, ROW_COUNT + 2):
        bridge[f'A{row_num}'] = f'=Accounts!A{row_num}'
        bridge[f'B{row_num}'] = f'=Accounts!B{row_num}'
        bridge[f'C{row_num}'] = f'=Accounts!C{row_num}-Accounts!E{row_num}'
        bridge[f'D{row_num}'] = f'=Accounts!C{row_num}*Accounts!F{row_num}'
        bridge[f'E{row_num}'] = f'=Accounts!D{row_num}*Assumptions!B3'
        bridge[f'F{row_num}'] = f'=D{row_num}+E{row_num}'
        bridge[f'G{row_num}'] = f'=Accounts!G{row_num}'
        bridge[f'H{row_num}'] = f'=F{row_num}*0'
        bridge[f'I{row_num}'] = f'=F{row_num}-G{row_num}'
        bridge[f'J{row_num}'] = f'=I{row_num}/0'

    wb.save(output_file)
    wb.close()


def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    output_file = os.path.join(args.output_path, 'arr_bridge.xlsx')
    build_workbook(output_file)
    with open(os.path.join(args.output_path, 'task_metadata.json'), 'w', encoding='utf-8') as f:
        json.dump(
            {
                'task_id': 'formula_repair/task_005_fix_assumption_driven_arr_bridge_formulas',
                'target_file': 'arr_bridge.xlsx',
            },
            f,
            ensure_ascii=False,
            indent=2,
        )
    print(os.path.abspath(args.output_path))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Create workspace for formula repair task')
    parser.add_argument('output_path', help='Directory to create workspace in')
    main(parser.parse_args())
