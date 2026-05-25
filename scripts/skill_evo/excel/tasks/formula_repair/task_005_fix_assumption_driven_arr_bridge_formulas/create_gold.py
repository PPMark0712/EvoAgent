import argparse
import os

from openpyxl import load_workbook

from create_workspace import ROW_COUNT, build_workbook


def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    target = os.path.join(args.output_path, 'arr_bridge.xlsx')
    if not os.path.isfile(target):
        build_workbook(target)
    wb = load_workbook(target)
    ws = wb['Bridge']
    for row_num in range(2, ROW_COUNT + 2):
        ws[f'C{row_num}'] = f'=Accounts!C{row_num}+Accounts!D{row_num}'
        ws[f'D{row_num}'] = f'=(Accounts!C{row_num}-Accounts!E{row_num})*Accounts!F{row_num}'
        ws[f'E{row_num}'] = f'=C{row_num}*Assumptions!$B$2*(1+Accounts!I{row_num})'
        ws[f'F{row_num}'] = f'=(D{row_num}+E{row_num})*(1-Assumptions!$B$3)'
        ws[f'G{row_num}'] = f'=Accounts!G{row_num}*Assumptions!$B$4'
        ws[f'H{row_num}'] = f'=F{row_num}*Accounts!H{row_num}'
        ws[f'I{row_num}'] = f'=F{row_num}-G{row_num}-H{row_num}+Accounts!J{row_num}+Assumptions!$B$5'
        ws[f'J{row_num}'] = f'=IF(F{row_num}=0,0,I{row_num}/F{row_num})'
    wb.save(target)
    wb.close()
    print(os.path.abspath(target))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Create gold workbook for formula repair task')
    parser.add_argument('output_path', help='Directory to place gold workbook in')
    main(parser.parse_args())
