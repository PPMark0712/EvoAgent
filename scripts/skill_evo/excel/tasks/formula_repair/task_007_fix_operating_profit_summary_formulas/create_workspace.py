import argparse
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Font


ROW_COUNT = 220
CHANNELS = ['Retail', 'Online', 'Partner', 'Field']
RATE_VALUES = {
    'Support Multiplier': 1.15,
    'Tax Rate': 0.24,
}


def canonical_rows():
    rows = []
    for idx in range(ROW_COUNT):
        gross_sales = 6800 + (idx % 20) * 240
        returns = 220 + (idx % 6) * 35
        discount_rate = 0.03 + (idx % 4) * 0.01
        orders = 55 + (idx % 11) * 3
        cogs_per_order = 34 + (idx % 5) * 4
        shipping_per_order = 8 + (idx % 4) * 1.5
        support_per_order = 5 + (idx % 3) * 1.2
        rows.append((
            f'STORE-{idx + 1:03d}',
            CHANNELS[idx % len(CHANNELS)],
            gross_sales,
            returns,
            discount_rate,
            orders,
            cogs_per_order,
            shipping_per_order,
            support_per_order,
            1,
        ))
    return rows


def build_workbook(output_file):
    wb = Workbook()
    transactions = wb.active
    transactions.title = 'Transactions'
    headers = [
        'Store ID',
        'Channel',
        'Gross Sales',
        'Returns',
        'Discount Rate',
        'Orders',
        'COGS per Order',
        'Shipping per Order',
        'Support per Order',
        'Active Flag',
    ]
    transactions.append(headers)
    for cell in transactions[1]:
        cell.font = Font(bold=True)
    for row in canonical_rows():
        transactions.append(list(row))

    rates = wb.create_sheet('Rates')
    rates['A1'] = 'Driver'
    rates['B1'] = 'Value'
    rates['A1'].font = Font(bold=True)
    rates['B1'].font = Font(bold=True)
    rates['A2'] = 'Support Multiplier'
    rates['B2'] = RATE_VALUES['Support Multiplier']
    rates['A3'] = 'Tax Rate'
    rates['B3'] = RATE_VALUES['Tax Rate']

    summary = wb.create_sheet('Summary')
    summary_headers = [
        'Store ID',
        'Channel',
        'Gross Revenue',
        'Discount Loss',
        'Net Revenue',
        'COGS',
        'Fulfillment Cost',
        'Support Cost',
        'Operating Profit',
        'Tax Expense',
        'Net Profit',
    ]
    summary.append(summary_headers)
    for cell in summary[1]:
        cell.font = Font(bold=True)
    for row_num in range(2, ROW_COUNT + 2):
        summary[f'A{row_num}'] = f'=Transactions!A{row_num}'
        summary[f'B{row_num}'] = f'=Transactions!B{row_num}'
        summary[f'C{row_num}'] = f'=Transactions!C{row_num}+Transactions!D{row_num}'
        summary[f'D{row_num}'] = f'=Transactions!C{row_num}*Transactions!E{row_num}'
        summary[f'E{row_num}'] = f'=C{row_num}'
        summary[f'F{row_num}'] = f'=Transactions!F{row_num}*Transactions!H{row_num}'
        summary[f'G{row_num}'] = f'=Transactions!F{row_num}*Transactions!G{row_num}'
        summary[f'H{row_num}'] = f'=Transactions!I{row_num}'
        summary[f'I{row_num}'] = f'=E{row_num}-F{row_num}'
        summary[f'J{row_num}'] = f'=I{row_num}*Rates!B2'
        summary[f'K{row_num}'] = f'=I{row_num}+J{row_num}'
    summary[f'A{ROW_COUNT + 2}'] = 'Network Total'

    wb.save(output_file)
    wb.close()


def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    output_file = os.path.join(args.output_path, 'operating_profit_summary.xlsx')
    build_workbook(output_file)
    with open(os.path.join(args.output_path, 'task_metadata.json'), 'w', encoding='utf-8') as f:
        json.dump(
            {
                'task_id': 'formula_repair/task_007_fix_operating_profit_summary_formulas',
                'target_file': 'operating_profit_summary.xlsx',
            },
            f,
            ensure_ascii=False,
            indent=2,
        )
    print(os.path.abspath(args.output_path))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Create workspace for formula repair task')
    parser.add_argument('output_path', help='Directory to create workspace in')
    main(parser.parse_args())
