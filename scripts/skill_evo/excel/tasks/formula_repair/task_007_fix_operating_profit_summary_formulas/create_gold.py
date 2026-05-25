import argparse
import os

from openpyxl import load_workbook
from openpyxl.styles import Font

from create_workspace import ROW_COUNT, build_workbook


TOTAL_ROW = ROW_COUNT + 2


def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    target = os.path.join(args.output_path, 'operating_profit_summary.xlsx')
    if not os.path.isfile(target):
        build_workbook(target)
    wb = load_workbook(target)
    ws = wb['Summary']
    for row_num in range(2, ROW_COUNT + 2):
        ws[f'C{row_num}'] = f'=Transactions!C{row_num}-Transactions!D{row_num}'
        ws[f'D{row_num}'] = f'=C{row_num}*Transactions!E{row_num}'
        ws[f'E{row_num}'] = f'=C{row_num}-D{row_num}'
        ws[f'F{row_num}'] = f'=Transactions!F{row_num}*Transactions!G{row_num}'
        ws[f'G{row_num}'] = f'=Transactions!F{row_num}*Transactions!H{row_num}'
        ws[f'H{row_num}'] = f'=Transactions!F{row_num}*Transactions!I{row_num}*Rates!$B$2'
        ws[f'I{row_num}'] = f'=E{row_num}-F{row_num}-G{row_num}-H{row_num}'
        ws[f'J{row_num}'] = f'=IF(I{row_num}<=0,0,I{row_num}*Rates!$B$3)'
        ws[f'K{row_num}'] = f'=I{row_num}-J{row_num}'
    ws[f'B{TOTAL_ROW}'] = 'All Channels'
    for coord in [f'A{TOTAL_ROW}', f'B{TOTAL_ROW}', f'C{TOTAL_ROW}', f'D{TOTAL_ROW}', f'E{TOTAL_ROW}', f'F{TOTAL_ROW}', f'G{TOTAL_ROW}', f'H{TOTAL_ROW}', f'I{TOTAL_ROW}', f'J{TOTAL_ROW}', f'K{TOTAL_ROW}']:
        ws[coord].font = Font(bold=True)
    for col in 'CDEFGHIJK':
        ws[f'{col}{TOTAL_ROW}'] = f'=SUM({col}2:{col}{ROW_COUNT + 1})'
    wb.save(target)
    wb.close()
    print(os.path.abspath(target))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Create gold workbook for formula repair task')
    parser.add_argument('output_path', help='Directory to place gold workbook in')
    main(parser.parse_args())
