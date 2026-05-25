import argparse
import os

from openpyxl import load_workbook


ROW_COUNT = 220


def build_workbook(output_file):
    from create_workspace import build_workbook as create_input_workbook

    create_input_workbook(output_file)


def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    target = os.path.join(args.output_path, "forecast_model.xlsx")
    if not os.path.isfile(target):
        build_workbook(target)
    wb = load_workbook(target)
    ws = wb["Forecast"]
    for row_num in range(2, ROW_COUNT + 2):
        ws[f"C{row_num}"] = f"=Inputs!C{row_num}*Inputs!D{row_num}"
        ws[f"D{row_num}"] = f"=Inputs!E{row_num}*Inputs!F{row_num}"
        ws[f"E{row_num}"] = f"=Inputs!G{row_num}*Inputs!H{row_num}"
        ws[f"F{row_num}"] = f"=Inputs!I{row_num}*Inputs!J{row_num}"
        ws[f"G{row_num}"] = f"=SUM(C{row_num}:F{row_num})"
    wb.save(target)
    wb.close()
    print(os.path.abspath(target))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Create gold workbook for formula repair task")
    parser.add_argument("output_path", help="Directory to place gold workbook in")
    main(parser.parse_args())
