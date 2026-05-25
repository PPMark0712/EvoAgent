import argparse
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Font


ROW_COUNT = 220
BLOCK_SIZE = 55
BLOCK_LABELS = ['North', 'South', 'East', 'West']
REORDER_STEP = 41


def canonical_rows():
    grouped_rows = []
    for idx in range(ROW_COUNT):
        region = BLOCK_LABELS[idx // BLOCK_SIZE]
        team_id = f'TEAM-{idx + 1:03d}'
        seats = 24 + (idx % 10) * 3
        occupancy_rate = round(0.55 + (idx % 6) * 0.05, 4)
        revenue_per_occ_seat = 860 + (idx % 7) * 65
        support_cost_per_seat = 140 + (idx % 5) * 18
        grouped_rows.append((team_id, region, seats, occupancy_rate, revenue_per_occ_seat, support_cost_per_seat))
    return [grouped_rows[(idx * REORDER_STEP) % ROW_COUNT] for idx in range(ROW_COUNT)]


def build_workbook(output_file):
    wb = Workbook()
    inputs = wb.active
    inputs.title = 'Staffing'
    inputs.append(['Team ID', 'Region', 'Seats', 'Occupancy Rate', 'Revenue Per Occupied Seat', 'Support Cost Per Seat'])
    for cell in inputs[1]:
        cell.font = Font(bold=True)
    for row in canonical_rows():
        inputs.append(list(row))

    detail = wb.create_sheet('Utilization')
    detail.append(['Team ID', 'Region', 'Occupied Seats', 'Revenue', 'Support Cost', 'Contribution', 'Utilization Rate'])
    for cell in detail[1]:
        cell.font = Font(bold=True)
    for row_num in range(2, ROW_COUNT + 2):
        detail[f'A{row_num}'] = f'=Staffing!A{row_num}'
        detail[f'B{row_num}'] = f'=Staffing!B{row_num}'
        detail[f'C{row_num}'] = f'=Staffing!C{row_num}*0'
        detail[f'D{row_num}'] = f'=Staffing!C{row_num}*Staffing!E{row_num}'
        detail[f'E{row_num}'] = f'=Staffing!C{row_num}*Staffing!D{row_num}'
        detail[f'F{row_num}'] = f'=D{row_num}+E{row_num}'
        detail[f'G{row_num}'] = f'=F{row_num}/0'

    summary = wb.create_sheet('Dashboard')
    summary.append(['Region', 'Total Revenue', 'Total Contribution', 'Average Utilization'])
    for cell in summary[1]:
        cell.font = Font(bold=True)
    for idx, label in enumerate(BLOCK_LABELS, start=2):
        summary[f'A{idx}'] = label
        start = 2 + (idx - 2) * BLOCK_SIZE
        end = start + BLOCK_SIZE - 1
        summary[f'B{idx}'] = f'=SUM(Utilization!C{start}:C{end})'
        summary[f'C{idx}'] = f'=SUM(Utilization!E{start}:E{end})'
        summary[f'D{idx}'] = f'=MAX(Utilization!G{start}:G{end})'

    wb.save(output_file)
    wb.close()


def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    output_file = os.path.join(args.output_path, 'capacity_dashboard_audit.xlsx')
    build_workbook(output_file)
    with open(os.path.join(args.output_path, 'task_metadata.json'), 'w', encoding='utf-8') as f:
        json.dump({'task_id': 'workbook_audit/task_003_audit_capacity_dashboard_workbook', 'target_file': 'capacity_dashboard_audit.xlsx'}, f, ensure_ascii=False, indent=2)
    print(os.path.abspath(args.output_path))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Create workspace for workbook audit task')
    parser.add_argument('output_path', help='Directory to create workspace in')
    main(parser.parse_args())
