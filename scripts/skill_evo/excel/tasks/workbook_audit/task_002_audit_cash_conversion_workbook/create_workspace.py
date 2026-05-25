import argparse
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Font


ROW_COUNT = 220
BLOCK_SIZE = 55
BLOCK_LABELS = ['Enterprise', 'Commercial', 'SMB', 'Public']
REORDER_STEP = 43


def canonical_rows():
    grouped_rows = []
    for idx in range(ROW_COUNT):
        region = BLOCK_LABELS[idx // BLOCK_SIZE]
        customer_id = f'CUST-{idx + 1:03d}'
        revenue = 5400 + idx * 43
        collection_rate = round(0.72 + (idx % 6) * 0.03, 4)
        deferred_rate = round(0.05 + (idx % 4) * 0.01, 4)
        payroll = 780 + (idx % 7) * 55
        vendor_spend = 260 + (idx % 5) * 40
        grouped_rows.append((customer_id, region, revenue, collection_rate, deferred_rate, payroll, vendor_spend))
    return [grouped_rows[(idx * REORDER_STEP) % ROW_COUNT] for idx in range(ROW_COUNT)]


def build_workbook(output_file):
    wb = Workbook()
    inputs = wb.active
    inputs.title = 'Inputs'
    inputs.append(['Customer ID', 'Segment', 'Invoiced Revenue', 'Collection Rate', 'Deferred Rate', 'Payroll', 'Vendor Spend'])
    for cell in inputs[1]:
        cell.font = Font(bold=True)
    for row in canonical_rows():
        inputs.append(list(row))

    detail = wb.create_sheet('CashFlow')
    detail.append(['Customer ID', 'Segment', 'Invoiced Revenue', 'Cash Collected', 'Deferred Revenue', 'Operating Spend', 'Net Cash', 'Cash Conversion'])
    for cell in detail[1]:
        cell.font = Font(bold=True)
    for row_num in range(2, ROW_COUNT + 2):
        detail[f'A{row_num}'] = f'=Inputs!A{row_num}'
        detail[f'B{row_num}'] = f'=Inputs!B{row_num}'
        detail[f'C{row_num}'] = f'=Inputs!C{row_num}*0'
        detail[f'D{row_num}'] = f'=C{row_num}*Inputs!E{row_num}'
        detail[f'E{row_num}'] = f'=C{row_num}*Inputs!D{row_num}'
        detail[f'F{row_num}'] = f'=Inputs!F{row_num}-Inputs!G{row_num}'
        detail[f'G{row_num}'] = f'=D{row_num}+F{row_num}'
        detail[f'H{row_num}'] = f'=G{row_num}/0'

    summary = wb.create_sheet('Covenant')
    summary.append(['Segment', 'Total Cash Collected', 'Total Net Cash', 'Average Cash Conversion'])
    for cell in summary[1]:
        cell.font = Font(bold=True)
    for idx, label in enumerate(BLOCK_LABELS, start=2):
        summary[f'A{idx}'] = label
        start = 2 + (idx - 2) * BLOCK_SIZE
        end = start + BLOCK_SIZE - 1
        summary[f'B{idx}'] = f'=SUM(CashFlow!C{start}:C{end})'
        summary[f'C{idx}'] = f'=SUM(CashFlow!E{start}:E{end})'
        summary[f'D{idx}'] = f'=MAX(CashFlow!H{start}:H{end})'

    wb.save(output_file)
    wb.close()


def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    output_file = os.path.join(args.output_path, 'cash_conversion_audit.xlsx')
    build_workbook(output_file)
    with open(os.path.join(args.output_path, 'task_metadata.json'), 'w', encoding='utf-8') as f:
        json.dump({'task_id': 'workbook_audit/task_002_audit_cash_conversion_workbook', 'target_file': 'cash_conversion_audit.xlsx'}, f, ensure_ascii=False, indent=2)
    print(os.path.abspath(args.output_path))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Create workspace for workbook audit task')
    parser.add_argument('output_path', help='Directory to create workspace in')
    main(parser.parse_args())
