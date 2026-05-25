import argparse
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Font


ROW_COUNT = 220
LOOKUP_STEP = 19


def canonical_rows():
    rows = []
    for idx in range(ROW_COUNT):
        sku = f'SKU-{idx + 1:03d}'
        region = ['North', 'South', 'East', 'West'][idx % 4]
        forecast_units = 140 + (idx % 15) * 9
        backlog_units = 12 + (idx % 6) * 4
        yield_rate = round(0.84 + (idx % 5) * 0.02, 4)
        unit_cost = 28 + (idx % 8) * 2
        expedite_rate = round(0.03 + (idx % 4) * 0.015, 4)
        rows.append((sku, region, forecast_units, backlog_units, yield_rate, unit_cost, expedite_rate))
    return rows


def reordered_rows(rows, step):
    return [rows[(idx * step) % len(rows)] for idx in range(len(rows))]


def build_workbook(output_file):
    wb = Workbook()
    rows = canonical_rows()
    supply_rows = reordered_rows(rows, LOOKUP_STEP)

    demand = wb.active
    demand.title = 'Demand'
    demand.append(['SKU', 'Region', 'Forecast Units', 'Backlog Units'])
    for cell in demand[1]:
        cell.font = Font(bold=True)

    supply = wb.create_sheet('Supply')
    supply.append(['SKU', 'Yield Rate', 'Unit Cost', 'Expedite Rate'])
    for cell in supply[1]:
        cell.font = Font(bold=True)

    plan = wb.create_sheet('Plan')
    plan.append(['SKU', 'Region', 'Demand Units', 'Production Units', 'Base Spend', 'Expedite Spend', 'Total Spend', 'Spend Per Demand Unit'])
    for cell in plan[1]:
        cell.font = Font(bold=True)

    for row in rows:
        sku, region, forecast_units, backlog_units, yield_rate, unit_cost, expedite_rate = row
        demand.append([sku, region, forecast_units, backlog_units])

    for row in supply_rows:
        sku, _region, _forecast_units, _backlog_units, yield_rate, unit_cost, expedite_rate = row
        supply.append([sku, yield_rate, unit_cost, expedite_rate])

    for row_num in range(2, ROW_COUNT + 2):
        plan[f'A{row_num}'] = f'=Demand!A{row_num}'
        plan[f'B{row_num}'] = f'=Demand!B{row_num}'
        plan[f'C{row_num}'] = f'=Demand!C{row_num}-Demand!D{row_num}'
        plan[f'D{row_num}'] = f'=Demand!C{row_num}*Supply!B{row_num}'
        plan[f'E{row_num}'] = f'=Demand!D{row_num}+Supply!C{row_num}'
        plan[f'F{row_num}'] = f'=E{row_num}/Supply!D{row_num}'
        plan[f'G{row_num}'] = f'=F{row_num}'
        plan[f'H{row_num}'] = f'=G{row_num}/0'

    wb.save(output_file)
    wb.close()


def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    output_file = os.path.join(args.output_path, 'inventory_network_plan.xlsx')
    build_workbook(output_file)
    with open(os.path.join(args.output_path, 'task_metadata.json'), 'w', encoding='utf-8') as f:
        json.dump({'task_id': 'cross_sheet_reasoning/task_002_repair_inventory_network_plan', 'target_file': 'inventory_network_plan.xlsx'}, f, ensure_ascii=False, indent=2)
    print(os.path.abspath(args.output_path))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Create workspace for cross sheet reasoning task')
    parser.add_argument('output_path', help='Directory to create workspace in')
    main(parser.parse_args())
