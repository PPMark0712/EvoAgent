import argparse
import os

from openpyxl import load_workbook
from create_workspace import ROW_COUNT, build_workbook


LOOKUP_RANGE_END = ROW_COUNT + 1


def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    target = os.path.join(args.output_path, 'hidden_helper_margin_stack.xlsx')
    if not os.path.isfile(target):
        build_workbook(target)
    wb = load_workbook(target)
    ws = wb['MarginStack']
    for row_num in range(2, ROW_COUNT + 2):
        ws[f'A{row_num}'] = f'=Orders!A{row_num}'
        ws[f'B{row_num}'] = f'=Orders!B{row_num}'
        ws[f'C{row_num}'] = f'=Orders!C{row_num}*SUMIF(PriceBook!$A$2:$A${LOOKUP_RANGE_END},$A{row_num},PriceBook!$B$2:$B${LOOKUP_RANGE_END})'
        ws[f'D{row_num}'] = f'=C{row_num}*SUMIF(PriceBook!$A$2:$A${LOOKUP_RANGE_END},$A{row_num},PriceBook!$C$2:$C${LOOKUP_RANGE_END})'
        ws[f'E{row_num}'] = f'=C{row_num}*SUMIF(Rebates!$A$2:$A${LOOKUP_RANGE_END},$A{row_num},Rebates!$B$2:$B${LOOKUP_RANGE_END})'
        ws[f'F{row_num}'] = f'=(C{row_num}-D{row_num}-E{row_num})*SUMIF(Terms!$A$2:$A${LOOKUP_RANGE_END},$A{row_num},Terms!$B$2:$B${LOOKUP_RANGE_END})'
        ws[f'G{row_num}'] = f'=F{row_num}*SUMIF(Rebates!$A$2:$A${LOOKUP_RANGE_END},$A{row_num},Rebates!$C$2:$C${LOOKUP_RANGE_END})'
        ws[f'H{row_num}'] = f'=F{row_num}+G{row_num}'
        ws[f'I{row_num}'] = f'=IF(H{row_num}=0,0,(H{row_num}-Orders!C{row_num}*SUMIF(Terms!$A$2:$A${LOOKUP_RANGE_END},$A{row_num},Terms!$C$2:$C${LOOKUP_RANGE_END}))/H{row_num})'
    wb.save(target)
    wb.close()
    print(os.path.abspath(target))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Create gold workbook for hard cross sheet reasoning task')
    parser.add_argument('output_path', help='Directory to place gold workbook in')
    main(parser.parse_args())
