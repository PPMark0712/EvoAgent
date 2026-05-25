import argparse
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Font


ROW_COUNT = 220
PRICEBOOK_STEP = 17
REBATE_STEP = 19
TERMS_STEP = 23


def canonical_rows():
    rows = []
    for idx in range(ROW_COUNT):
        order_id = f'ORD-{idx + 1:03d}'
        region = ['North', 'South', 'East', 'West'][idx % 4]
        units = 48 + (idx % 13) * 4
        asp = 205 + (idx % 11) * 9
        discount_rate = round(0.03 + (idx % 5) * 0.0125, 4)
        rebate_rate = round(0.015 + (idx % 4) * 0.01, 4)
        service_attach = round(0.05 + (idx % 5) * 0.02, 4)
        renewal_rate = round(0.72 + (idx % 6) * 0.04, 4)
        support_cost_per_unit = 88 + (idx % 7) * 5
        rows.append((order_id, region, units, asp, discount_rate, rebate_rate, service_attach, renewal_rate, support_cost_per_unit))
    return rows


def reordered_rows(rows, step):
    return [rows[(idx * step) % len(rows)] for idx in range(len(rows))]


def build_workbook(output_file):
    wb = Workbook()
    rows = canonical_rows()
    pricebook_rows = reordered_rows(rows, PRICEBOOK_STEP)
    rebate_rows = reordered_rows(rows, REBATE_STEP)
    term_rows = reordered_rows(rows, TERMS_STEP)

    orders = wb.active
    orders.title = 'Orders'
    orders.append(['Order ID', 'Region', 'Units'])
    for cell in orders[1]:
        cell.font = Font(bold=True)

    pricebook = wb.create_sheet('PriceBook')
    pricebook.append(['Order ID', 'ASP', 'Discount Rate'])
    for cell in pricebook[1]:
        cell.font = Font(bold=True)

    rebates = wb.create_sheet('Rebates')
    rebates.append(['Order ID', 'Rebate Rate', 'Service Attach'])
    for cell in rebates[1]:
        cell.font = Font(bold=True)

    terms = wb.create_sheet('Terms')
    terms.append(['Order ID', 'Renewal Rate', 'Support Cost Per Unit'])
    for cell in terms[1]:
        cell.font = Font(bold=True)
    terms.sheet_state = 'hidden'

    output = wb.create_sheet('MarginStack')
    output.append([
        'Order ID', 'Region', 'Gross Revenue', 'Discount Loss', 'Rebate Loss',
        'Net Revenue', 'Services ARR', 'Final ARR', 'Final Margin'
    ])
    for cell in output[1]:
        cell.font = Font(bold=True)

    for order_id, region, units, asp, discount_rate, rebate_rate, service_attach, renewal_rate, support_cost_per_unit in rows:
        orders.append([order_id, region, units])

    for order_id, _region, _units, asp, discount_rate, _rebate_rate, _service_attach, _renewal_rate, _support_cost_per_unit in pricebook_rows:
        pricebook.append([order_id, asp, discount_rate])

    for order_id, _region, _units, _asp, _discount_rate, rebate_rate, service_attach, _renewal_rate, _support_cost_per_unit in rebate_rows:
        rebates.append([order_id, rebate_rate, service_attach])

    for order_id, _region, _units, _asp, _discount_rate, _rebate_rate, _service_attach, renewal_rate, support_cost_per_unit in term_rows:
        terms.append([order_id, renewal_rate, support_cost_per_unit])

    for row_num in range(2, ROW_COUNT + 2):
        output[f'A{row_num}'] = f'=Orders!A{row_num}'
        output[f'B{row_num}'] = f'=Orders!B{row_num}'
        output[f'C{row_num}'] = f'=Orders!C{row_num}+PriceBook!B{row_num}'
        output[f'D{row_num}'] = f'=C{row_num}*Terms!C{row_num}'
        output[f'E{row_num}'] = f'=C{row_num}*Rebates!C{row_num}'
        output[f'F{row_num}'] = f'=C{row_num}-D{row_num}-E{row_num}'
        output[f'G{row_num}'] = f'=F{row_num}*Terms!B{row_num}'
        output[f'H{row_num}'] = f'=F{row_num}'
        output[f'I{row_num}'] = f'=H{row_num}/0'

    wb.save(output_file)
    wb.close()


def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    output_file = os.path.join(args.output_path, 'hidden_helper_margin_stack.xlsx')
    build_workbook(output_file)
    with open(os.path.join(args.output_path, 'task_metadata.json'), 'w', encoding='utf-8') as f:
        json.dump({'task_id': 'cross_sheet_reasoning/task_004_repair_hidden_helper_margin_stack', 'target_file': 'hidden_helper_margin_stack.xlsx'}, f, ensure_ascii=False, indent=2)
    print(os.path.abspath(args.output_path))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Create workspace for hard cross sheet reasoning task')
    parser.add_argument('output_path', help='Directory to create workspace in')
    main(parser.parse_args())
