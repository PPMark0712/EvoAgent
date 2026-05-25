import argparse
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Font


ROW_COUNT = 220
PRICING_STEP = 17
COST_STEP = 19


def canonical_rows():
    rows = []
    for idx in range(ROW_COUNT):
        order_id = f'ORD-{idx + 1:03d}'
        region = ['North', 'South', 'East', 'West'][idx % 4]
        units = 45 + (idx % 14) * 3
        asp = 210 + (idx % 9) * 11
        discount_rate = round(0.03 + (idx % 5) * 0.015, 4)
        unit_cost = 118 + (idx % 7) * 6
        support_cost = 85 + (idx % 6) * 12
        rows.append((order_id, region, units, asp, discount_rate, unit_cost, support_cost))
    return rows


def reordered_rows(rows, step):
    return [rows[(idx * step) % len(rows)] for idx in range(len(rows))]


def build_workbook(output_file):
    wb = Workbook()
    rows = canonical_rows()
    pricing_rows = reordered_rows(rows, PRICING_STEP)
    cost_rows = reordered_rows(rows, COST_STEP)

    orders = wb.active
    orders.title = 'Orders'
    orders.append(['Order ID', 'Region', 'Units'])
    for cell in orders[1]:
        cell.font = Font(bold=True)

    pricing = wb.create_sheet('Pricing')
    pricing.append(['Order ID', 'ASP', 'Discount Rate'])
    for cell in pricing[1]:
        cell.font = Font(bold=True)

    costs = wb.create_sheet('Costs')
    costs.append(['Order ID', 'Unit Cost', 'Support Cost'])
    for cell in costs[1]:
        cell.font = Font(bold=True)

    bridge = wb.create_sheet('Bridge')
    bridge.append(['Order ID', 'Region', 'Gross Revenue', 'Discount Loss', 'Net Revenue', 'COGS', 'Contribution', 'Contribution Margin'])
    for cell in bridge[1]:
        cell.font = Font(bold=True)

    for row in rows:
        order_id, region, units, asp, discount_rate, unit_cost, support_cost = row
        orders.append([order_id, region, units])

    for row in pricing_rows:
        order_id, _region, _units, asp, discount_rate, _unit_cost, _support_cost = row
        pricing.append([order_id, asp, discount_rate])

    for row in cost_rows:
        order_id, _region, _units, _asp, _discount_rate, unit_cost, support_cost = row
        costs.append([order_id, unit_cost, support_cost])

    for row_num in range(2, ROW_COUNT + 2):
        bridge[f'A{row_num}'] = f'=Orders!A{row_num}'
        bridge[f'B{row_num}'] = f'=Orders!B{row_num}'
        bridge[f'C{row_num}'] = f'=Orders!C{row_num}+Pricing!B{row_num}'
        bridge[f'D{row_num}'] = f'=C{row_num}*Costs!C{row_num}'
        bridge[f'E{row_num}'] = f'=C{row_num}+D{row_num}'
        bridge[f'F{row_num}'] = f'=Costs!B{row_num}'
        bridge[f'G{row_num}'] = f'=E{row_num}-F{row_num}'
        bridge[f'H{row_num}'] = f'=G{row_num}/0'

    wb.save(output_file)
    wb.close()


def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    output_file = os.path.join(args.output_path, 'multi_sheet_margin_bridge.xlsx')
    build_workbook(output_file)
    with open(os.path.join(args.output_path, 'task_metadata.json'), 'w', encoding='utf-8') as f:
        json.dump({'task_id': 'cross_sheet_reasoning/task_003_repair_multi_sheet_margin_bridge', 'target_file': 'multi_sheet_margin_bridge.xlsx'}, f, ensure_ascii=False, indent=2)
    print(os.path.abspath(args.output_path))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Create workspace for cross sheet reasoning task')
    parser.add_argument('output_path', help='Directory to create workspace in')
    main(parser.parse_args())
