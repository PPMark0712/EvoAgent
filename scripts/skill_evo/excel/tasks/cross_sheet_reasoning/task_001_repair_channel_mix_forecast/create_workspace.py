import argparse
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Font


ROW_COUNT = 220
LOOKUP_STEP = 17
MISSING_DISCOUNT_STEP = 14
MISSING_ATTACH_STEP = 19


def canonical_rows():
    rows = []
    for idx in range(ROW_COUNT):
        account_id = f'ACCT-{idx + 1:03d}'
        region = ['North', 'South', 'East', 'West'][idx % 4]
        base_units = 85 + (idx % 12) * 4
        base_price = 120 + (idx % 9) * 6
        renewal_probability = round(0.62 + (idx % 5) * 0.05, 4)
        enterprise_share = round(0.35 + (idx % 4) * 0.08, 4)
        partner_discount = round(0.04 + (idx % 5) * 0.01, 4)
        service_attach = round(0.06 + (idx % 4) * 0.02, 4)
        rows.append((account_id, region, base_units, base_price, renewal_probability, enterprise_share, partner_discount, service_attach))
    return rows


def reordered_rows(rows, step):
    return [rows[(idx * step) % len(rows)] for idx in range(len(rows))]


def build_mix_rows():
    rows = []
    for idx, row in enumerate(reordered_rows(canonical_rows(), LOOKUP_STEP), start=1):
        account_id, region, _base_units, _base_price, _renewal_probability, enterprise_share, partner_discount, service_attach = row
        active_discount = partner_discount if idx % MISSING_DISCOUNT_STEP else None
        active_attach = service_attach if idx % MISSING_ATTACH_STEP else None
        rows.append((account_id, region, 'Active', enterprise_share, active_discount, active_attach))
        if idx % 9 == 0:
            rows.append((account_id, region, 'Archived', round(min(0.95, enterprise_share + 0.18), 4), 0.22, 0.18))
        if idx % 13 == 0:
            distractor_region = ['North', 'South', 'East', 'West'][(idx + 1) % 4]
            rows.append((account_id, distractor_region, 'Archived', 0.15, 0.09, 0.05))
    rows.append(('NOTE', 'N/A', 'Ignore', None, None, None))
    rows.append((None, None, None, None, None, None))
    return rows


def build_workbook(output_file):
    wb = Workbook()
    rows = canonical_rows()
    mix_rows = build_mix_rows()

    drivers = wb.active
    drivers.title = 'Drivers'
    drivers.append(['Account ID', 'Region', 'Base Units', 'Base Price', 'Renewal Probability'])
    for cell in drivers[1]:
        cell.font = Font(bold=True)

    mix = wb.create_sheet('Mix')
    mix.append(['Account ID', 'Region', 'Status', 'Enterprise Share', 'Partner Discount', 'Service Attach'])
    for cell in mix[1]:
        cell.font = Font(bold=True)

    output = wb.create_sheet('Output')
    output.append(['Account ID', 'Region', 'Gross Sales', 'Enterprise Sales', 'Partner Discount', 'Net Sales', 'Services', 'Total ARR'])
    for cell in output[1]:
        cell.font = Font(bold=True)

    for row in rows:
        account_id, region, base_units, base_price, renewal_probability, enterprise_share, partner_discount, service_attach = row
        drivers.append([account_id, region, base_units, base_price, renewal_probability])

    for row in mix_rows:
        mix.append(list(row))

    for row_num in range(2, ROW_COUNT + 2):
        output[f'A{row_num}'] = f'=Drivers!A{row_num}'
        output[f'B{row_num}'] = f'=Drivers!B{row_num}'
        output[f'C{row_num}'] = f'=Drivers!C{row_num}*Drivers!D{row_num}'
        output[f'D{row_num}'] = f'=C{row_num}*Mix!D{row_num}'
        output[f'E{row_num}'] = f'=C{row_num}*Mix!E{row_num}'
        output[f'F{row_num}'] = f'=C{row_num}-E{row_num}'
        output[f'G{row_num}'] = f'=F{row_num}*Mix!F{row_num}'
        output[f'H{row_num}'] = f'=F{row_num}'

    wb.save(output_file)
    wb.close()


def main(args):
    os.makedirs(args.output_path, exist_ok=True)
    output_file = os.path.join(args.output_path, 'channel_mix_forecast.xlsx')
    build_workbook(output_file)
    with open(os.path.join(args.output_path, 'task_metadata.json'), 'w', encoding='utf-8') as f:
        json.dump({'task_id': 'cross_sheet_reasoning/task_001_repair_channel_mix_forecast', 'target_file': 'channel_mix_forecast.xlsx'}, f, ensure_ascii=False, indent=2)
    print(os.path.abspath(args.output_path))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Create workspace for cross sheet reasoning task')
    parser.add_argument('output_path', help='Directory to create workspace in')
    main(parser.parse_args())
